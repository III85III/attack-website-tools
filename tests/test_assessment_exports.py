"""Regression coverage for assessment identity, grouping, and complete exports."""
import json

import pytest
from docx import Document

from src.core.database import DatabaseManager
from src.core.models import Finding, FindingConfidence, FindingSeverity, ScanRun
from src.reporting.export_service import export_assessment


def finding(scan, title):
    return Finding(scan_id=scan.id, title=title, description="Controlled test evidence",
                   source_tool=scan.tool_name, target=scan.target, severity=FindingSeverity.HIGH,
                   severity_source="Fixture", confidence=FindingConfidence.HIGH,
                   confidence_justification="Fixture", evidence_raw="Fixture response")


def test_assessment_reports_include_child_engines_and_exclude_other_targets(tmp_path):
    db = DatabaseManager(tmp_path / "exports.db")
    try:
        parent = ScanRun(tool_name="Sentinel DAST", target="https://first.test", raw_command="", status="COMPLETED")
        child = ScanRun(tool_name="Nmap", target=parent.target, raw_command="", parent_scan_id=parent.id, status="COMPLETED")
        other = ScanRun(tool_name="Sentinel DAST", target="https://unrelated.test", raw_command="", status="COMPLETED")
        for scan in (parent, child, other):
            db.save_scan(scan)
        a, b, excluded = finding(parent, "Parent evidence"), finding(child, "Child engine evidence"), finding(other, "Unrelated evidence")
        for item in (a, b, excluded):
            db.save_finding(item)
        # Reproduce worker finalization after findings have already been stored.
        parent.findings_count = 2
        db.save_scan(parent)
        assert {s.id for s in db.list_scans(root_only=True)} == {parent.id, other.id}
        assert db.get_scan(child.id).parent_scan_id == parent.id
        report = export_assessment(db, parent.id, "json", tmp_path)
        payload = json.loads(report.read_text(encoding="utf-8"))
        assert {f["id"] for f in payload["findings"]} == {a.id, b.id}
        assert payload["scan"]["id"] == parent.id
        docx = export_assessment(db, parent.id, "docx", tmp_path)
        document = Document(docx)
        text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
        assert a.title in text and b.title in text
        assert excluded.title not in text and other.target not in text
    finally:
        db.close()


def test_export_has_no_ui_row_limit_and_does_not_overwrite_previous_file(tmp_path):
    db = DatabaseManager(tmp_path / "all-findings.db")
    try:
        scan = ScanRun(tool_name="Sentinel", target="https://first.test", raw_command="", status="COMPLETED")
        db.save_scan(scan)
        for i in range(205):
            db.save_finding(finding(scan, f"Finding {i}"))
        assert len(db.list_findings(scan_id=scan.id)) == 200
        first = export_assessment(db, scan.id, "json", tmp_path)
        second = export_assessment(db, scan.id, "json", tmp_path)
        assert first != second
        assert json.loads(first.read_text(encoding="utf-8"))["total_findings"] == 205
    finally:
        db.close()


def test_running_or_missing_assessment_cannot_be_exported(tmp_path):
    db = DatabaseManager(tmp_path / "pending.db")
    try:
        scan = ScanRun(tool_name="Sentinel", target="https://first.test", raw_command="")
        db.save_scan(scan)
        with pytest.raises(ValueError, match="finish"):
            export_assessment(db, scan.id, "json", tmp_path)
        with pytest.raises(ValueError, match="no longer exists"):
            export_assessment(db, "missing", "docx", tmp_path)
    finally:
        db.close()
