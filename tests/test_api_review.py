import json

import pytest

from src.adapters.api_review import APIReviewAdapter, build_plan, import_spec_file, validate_api_plan
from src.core.models import ScanRun, FindingStatus, ToolHealthInfo
from test_browser_xss import local_target


def specification(paths=None):
    return {"openapi": "3.0.4", "info": {"title": "Fixture", "version": "1"}, "servers": [{"url": "/"}],
            "paths": paths or {"/reflect": {"get": {"parameters": [{"in": "query", "name": "q", "schema": {"type": "string", "example": "PRIVATE_EXAMPLE"}}]}}}}


def plan(target, document=None):
    return build_plan(target, document or specification(), "a" * 64)


def run(target, config):
    return APIReviewAdapter().run_scan(target, ScanRun(tool_name=APIReviewAdapter.name, target=target, raw_command=""), config)


def test_schema_supplies_real_unlinked_xss_input(local_target, tmp_path):
    from src.reporting.docx_report import ReportGenerator
    from docx import Document
    target, observed = local_target
    document = specification({"/attribute": specification()["paths"]["/reflect"], "/items/{id}": {"get": {}}, "/write": {"post": {}}})
    config = plan(target, document)
    scan, findings = run(target, config)
    assert scan.status == "PARTIAL" and scan.request_count == 3
    assert len(findings) == 1 and findings[0].status == FindingStatus.VALIDATED
    coverage = scan.execution_log[0]
    assert coverage["browser_request_count"] == 1
    assert scan.request_count + coverage["browser_request_count"] == len(observed)
    assert len(coverage["checks"]) == 3 and coverage["checks"][0]["outcome"] == "COLLECTED"
    assert coverage["checks"][1]["outcome"] == coverage["checks"][2]["outcome"] == "NOT_TESTED"
    output = ReportGenerator(tmp_path).generate_docx(scan, findings, [])
    doc = Document(output)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "پوشش مسیرهای API" in text and "NOT_TESTED" in text and "پارامتر آسیب‌پذیر" in text
    assert "PRIVATE_EXAMPLE" not in json.dumps(config) + text


@pytest.mark.parametrize("path", ["/escaped", "/csp"])
def test_api_does_not_validate_blocked_or_encoded_xss(local_target, path):
    target, _ = local_target
    scan, findings = run(target, plan(target, specification({path: specification()["paths"]["/reflect"]})))
    assert scan.status == "COMPLETED"
    assert not any(item.status == FindingStatus.VALIDATED for item in findings)


def test_non_success_response_is_inconclusive_and_redirect_is_not_followed(local_target):
    target, observed = local_target
    scan, findings = run(target, plan(target, specification({"/redirect": {"get": {}}})))
    assert scan.status == "PARTIAL" and not findings
    assert scan.execution_log[0]["checks"][0]["outcome"] == "INCONCLUSIVE"
    assert len(observed) == 1


@pytest.mark.parametrize("server", ["https://outside.invalid", "//outside.invalid/api", "http://user:secret@127.0.0.1/", "/api?token=secret", "/api/{tenant}", "/../api"])
def test_unsafe_servers_rejected_before_traffic(server):
    document = specification()
    document["servers"] = [{"url": server}]
    with pytest.raises(ValueError):
        plan("http://127.0.0.1/", document)


def test_swagger_basepath_and_security_metadata():
    config = plan("https://owned.invalid/", {"swagger": "2.0", "host": "owned.invalid", "basePath": "/api/v2", "schemes": ["https"],
                                           "security": [{"session": []}], "paths": {"/status": {"get": {}}}})
    assert config["operations"][0]["path"] == "/api/v2/status"
    assert config["operations"][0]["security_declared"]


def test_auth_declaration_mismatch_is_potential_and_empty_alternative_allows_public(local_target):
    target, _ = local_target
    document = specification({"/public": {"get": {}}, "/optional": {"get": {"security": [{}]}}})
    document["security"] = [{"bearer": []}]
    scan, findings = run(target, plan(target, document))
    # The fixture has empty responses without query input, so add a query parameter.
    assert not findings
    document["paths"]["/public"]["get"]["parameters"] = [{"in": "query", "name": "q"}]
    document["paths"]["/optional"]["get"]["parameters"] = [{"in": "query", "name": "q"}]
    # Avoid browser work here: the encoded path exercises only declaration checks.
    document["paths"]["/escaped"] = document["paths"].pop("/public")
    document["paths"].pop("/optional")
    scan, findings = run(target, plan(target, document))
    assert len(findings) == 1 and findings[0].status == FindingStatus.POTENTIAL and findings[0].cwe_id == "CWE-306"
    document["paths"]["/escaped"]["get"]["security"] = [{}]
    assert not run(target, plan(target, document))[1]


def test_each_api_get_and_probe_drops_previous_response_cookies(local_target):
    target, observed = local_target
    document = specification({"/give-cookie": {"get": {}}, "/escaped": specification()["paths"]["/reflect"]})
    scan, _ = run(target, plan(target, document))
    assert scan.request_count == len(observed) == 4
    assert all(cookie is None for _, _, cookie in observed)


def test_all_exclusions_are_recorded_and_no_references_fetched(local_target):
    target, observed = local_target
    paths = {"/write": {"post": {}}, "/logout": {"get": {}}, "/item/{id}": {"get": {}},
             "/reference": {"get": {"parameters": [{"$ref": "https://outside.invalid/schema"}]}},
             "/credential": {"get": {"parameters": [{"name": "api_key", "in": "query"}]}}}
    scan, findings = run(target, plan(target, specification(paths)))
    assert scan.status == "PARTIAL" and scan.request_count == 0 and not findings and not observed
    assert len(scan.execution_log[0]["checks"]) == 5
    assert all(item["reason"] for item in scan.execution_log[0]["checks"])


def test_changed_target_and_tampered_method_rejected_before_traffic(local_target):
    target, observed = local_target
    config = plan(target)
    with pytest.raises(ValueError):
        validate_api_plan("https://outside.invalid", config)
    config["operations"][0]["method"] = "POST"
    scan, _ = run(target, config)
    assert scan.status == "FAILED" and not observed


def test_spec_size_duplicates_and_invalid_parameter_rejected(tmp_path):
    source = tmp_path / "api.json"
    source.write_text('{"openapi":"3.0.4","openapi":"3.1.0"}')
    with pytest.raises(ValueError):
        import_spec_file("https://owned.invalid", source)
    source.write_bytes(b" " * 2_000_001)
    with pytest.raises(ValueError):
        import_spec_file("https://owned.invalid", source)
    with pytest.raises(ValueError):
        plan("https://owned.invalid", specification({"/x": {"get": {"parameters": [None]}}}))


def test_local_api_diagnostic_exports_actual_word(local_target, tmp_path):
    from main import _verify_local_api
    target, observed = local_target
    source = tmp_path / "api.json"
    source.write_text(json.dumps(specification({"/escaped": specification()["paths"]["/reflect"]})))
    output = tmp_path / "result.json"
    assert _verify_local_api(target, str(source), str(output)) == 0
    value = json.loads(output.read_text())
    assert value["word_export_verified"] and value["export_verified"] and value["persistence_verified"]
    assert value["scan"]["request_count"] == len(observed) == 3
    assert _verify_local_api("https://outside.invalid", str(source), str(output)) == 1
    assert len(observed) == 3


def test_operation_limit_is_reported_without_silent_truncation(local_target, monkeypatch):
    target, observed = local_target
    monkeypatch.setattr("src.adapters.api_review.MAX_API_GETS", 2)
    scan, _ = run(target, plan(target, specification({f"/p{i}": {"get": {}} for i in range(4)})))
    assert scan.request_count == len(observed) == 2 and scan.status == "PARTIAL"
    assert len(scan.execution_log[0]["checks"]) == 4


def test_worker_links_api_child_and_exports_its_coverage(local_target, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QApplication
    from src.ui.tabs import WebAuditWorker, WebSecurityTab
    from src.core.database import DatabaseManager
    from src.reporting.export_service import export_assessment
    app = QApplication.instance() or QApplication([])
    target, _ = local_target
    database = DatabaseManager(tmp_path / "api.db")
    monkeypatch.setattr(DatabaseManager, "get_instance", classmethod(lambda cls: database))
    for adapter in ("ZapAdapter", "NmapAdapter"):
        monkeypatch.setattr("src.ui.tabs." + adapter + ".detect", lambda self: ToolHealthInfo(name="Fixture", installed=False))
    monkeypatch.setattr("src.ui.tabs.write_execution_log", lambda scan: tmp_path / "log.json")
    worker = WebAuditWorker(target, api_plan=plan(target, specification({"/escaped": specification()["paths"]["/reflect"]})))
    results = []
    worker.finished_signal.connect(lambda scan, findings: results.append((scan, findings)))
    worker.run()
    assert results[0][0].status == "COMPLETED" and worker.api_plan is None
    children = [item for item in database.list_scans() if item.parent_scan_id == results[0][0].id]
    assert len(children) == 1 and children[0].request_count == 3
    payload = json.loads(export_assessment(database, results[0][0].id, "json", tmp_path).read_text())
    assert any(item.get("operation") == "API coverage" for item in payload["scan"]["execution_log"])
    widget = WebSecurityTab()
    widget.api_plan = plan(target)
    widget.clear_api_definition()
    assert widget.api_plan is None and widget.session_btn.menu() is not None
    widget.close()
    database.close()


def test_malformed_security_nonfinite_json_and_inventory_overflow_rejected(tmp_path):
    source = tmp_path / "invalid.json"
    source.write_text('{"openapi":"3.0.4","unexpected":NaN}')
    with pytest.raises(ValueError):
        import_spec_file("https://owned.invalid", source)
    document = specification()
    document["security"] = {"not": "an array"}
    with pytest.raises(ValueError):
        plan("https://owned.invalid", document)
    with pytest.raises(ValueError):
        plan("https://owned.invalid", {"swagger": "2.0", "host": 123, "paths": {"/": {"get": {}}}})
    with pytest.raises(ValueError):
        plan("https://owned.invalid", specification({f"/p{i}": {"get": {}, "post": {}, "put": {}} for i in range(200)}))


def test_oversized_version_is_rejected():
    document = specification()
    document["openapi"] = "3.0." + "9" * 100000
    with pytest.raises(ValueError):
        plan("https://owned.invalid", document)


def test_reference_only_path_remains_visible_without_fetching(local_target):
    target, observed = local_target
    config = plan(target, specification({"/external": {"$ref": "https://outside.invalid/private"}, "/local": {"$ref": "#/components/pathItems/read"}}))
    scan, findings = run(target, config)
    assert scan.status == "PARTIAL" and scan.request_count == 0 and not findings and not observed
    assert len(scan.execution_log[0]["checks"]) == 2
    assert all(item["method"] == "REFERENCE" and item["outcome"] == "NOT_TESTED" for item in scan.execution_log[0]["checks"])
