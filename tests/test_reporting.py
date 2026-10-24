"""
Unit tests for DOCX and JSON report generators.
"""
import json
import tempfile
import pytest
from pathlib import Path

from src.core.models import (
    Finding,
    FindingConfidence,
    FindingSeverity,
    FindingStatus,
    ScanRun,
    ToolHealthInfo,
)
from src.reporting.docx_report import JSONExporter, ReportGenerator


@pytest.mark.parametrize("status", [FindingStatus.SCANNER_REPORTED, FindingStatus.VALIDATED])
def test_reflection_evidence_never_claims_browser_execution(tmp_path, status):
    from docx import Document
    item = Finding(title="Potential unescaped reflected input", description="Inert marker reflection fixture",
                   source_tool="Web Security Auditor", target="https://owned.invalid/search", cwe_id="CWE-79",
                   severity=FindingSeverity.MEDIUM, severity_source="Reflection", confidence=FindingConfidence.MEDIUM,
                   confidence_justification="No browser execution", status=status,
                   evidence_raw=json.dumps({"verification_method": "inert_marker_reflection", "script_execution_verified": False}))
    path = ReportGenerator(tmp_path).generate_docx(None, [item], [])
    document = Document(path)
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for table in document.tables for row in table.rows for c in row.cells])
    assert "تفکیک وضعیت شواهد" in text
    assert "اجرای کد در مرورگر تأیید نشده" in text


@pytest.mark.parametrize("matching", [True, False])
def test_browser_proof_report_requires_matching_observation(tmp_path, matching):
    from docx import Document
    proof = {"script_execution_verified": True, "proof_expected": "fixture-123",
             "proof_observed": "fixture-123" if matching else "different", "browser_name": "FixtureBrowser",
             "browser_version": "1", "parameter": "q", "payload": "BENIGN_FIXTURE"}
    item = Finding(title="Browser-confirmed reflected XSS", description="Browser proof fixture",
                   source_tool="Web Security Auditor", target="https://owned.invalid/search", cwe_id="CWE-79",
                   severity=FindingSeverity.HIGH, severity_source="Execution", confidence=FindingConfidence.CONFIRMED,
                   confidence_justification="Observed", status=FindingStatus.VALIDATED,
                   evidence_raw=json.dumps({"verification_method": "browser_execution_probe", "script_execution_verified": True,
                                            "browser_verification": proof}))
    path = ReportGenerator(tmp_path).generate_docx(None, [item], [])
    text = "\n".join(c.text for table in Document(path).tables for row in table.rows for c in row.cells)
    assert ("مقدار محاسبه‌شده با مقدار مورد انتظار برابر بود" in text) is matching
    if matching:
        assert "BENIGN_FIXTURE" in text and "پارامتر آسیب‌پذیر" in text
    else:
        assert "شواهد کافی برای تأیید اجرای کد" in text


def test_docx_report_generation():
    with tempfile.TemporaryDirectory() as tmpdir:
        gen = ReportGenerator(output_dir=Path(tmpdir))

        scan = ScanRun(tool_name="Nmap", target="127.0.0.1", raw_command="nmap 127.0.0.1")
        finding = Finding(
            title="Open SMB Port 445",
            description="SMB Port 445 is listening.",
            source_tool="Nmap",
            target="127.0.0.1:445",
            severity=FindingSeverity.MEDIUM,
            severity_source="Nmap",
            confidence=FindingConfidence.CONFIRMED,
            confidence_justification="syn-ack",
            evidence_raw="port 445 open",
        )
        tool = ToolHealthInfo(name="Nmap", installed=True, version="7.99")

        docx_path = gen.generate_docx(scan, [finding], [tool], "127.0.0.1")
        assert docx_path.is_file()
        assert docx_path.stat().st_size > 1000  # Valid docx document


def test_json_export_generation():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = JSONExporter(output_dir=Path(tmpdir))

        scan = ScanRun(tool_name="Nmap", target="127.0.0.1", raw_command="nmap 127.0.0.1")
        finding = Finding(
            title="Open SMB Port 445",
            description="SMB Port 445 is listening.",
            source_tool="Nmap",
            target="127.0.0.1:445",
            severity=FindingSeverity.MEDIUM,
            severity_source="Nmap",
            confidence=FindingConfidence.CONFIRMED,
            confidence_justification="syn-ack",
            evidence_raw="port 445 open",
        )
        tool = ToolHealthInfo(name="Nmap", installed=True, version="7.99")

        json_path = exporter.export_json(scan, [finding], [tool])
        assert json_path.is_file()

        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            assert data["platform"] == "III85III QOLAM SENTINEL"
            assert data["total_findings"] == 1
            assert len(data["findings"]) == 1
