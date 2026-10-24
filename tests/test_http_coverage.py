import json

import pytest
from docx import Document

from test_response_collection import response, run
from src.adapters.web_auditor import WebAuditorAdapter
from src.reporting.docx_report import JSONExporter, ReportGenerator


@pytest.mark.parametrize("status", [301, 302, 401, 403, 404, 429, 500, 503])
def test_initial_non_success_is_partial_and_reported(tmp_path, monkeypatch, status):
    adapter = WebAuditorAdapter()
    calls = []
    def get(url, **kwargs):
        assert kwargs["allow_redirects"] is False
        assert url.startswith("http://127.0.0.1/")
        calls.append(url)
        value = response()
        value.status_code = status
        value.headers["Location"] = "https://not-contacted.invalid/?token=UNFOLLOWED_SECRET"
        return value
    monkeypatch.setattr(adapter.session, "get", get)
    scan, findings = run(adapter)
    assert scan.status == "PARTIAL" and scan.response_status == status
    assert scan.request_count == len(calls) > 0
    assert scan.execution_log[-1]["initial_response_coverage_limited"] is True
    assert "coverage was not established" in scan.completion_reason or "destination content was not assessed" in scan.completion_reason
    assert "UNFOLLOWED_SECRET" not in json.dumps(scan.to_dict())
    payload = json.loads(JSONExporter(tmp_path).export_json(scan, findings, []).read_text(encoding="utf-8"))
    assert payload["scan"]["status"] == "PARTIAL"
    assert payload["scan"]["response_status"] == status
    word = Document(ReportGenerator(tmp_path).generate_docx(scan, findings, []))
    text = "\n".join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
    assert scan.execution_log[-1]["initial_response_limitation"] in text


@pytest.mark.parametrize("status", [200, 201, 204])
def test_success_response_is_not_marked_partial_by_status_alone(monkeypatch, status):
    adapter = WebAuditorAdapter()
    def get(*args, **kwargs):
        value = response()
        value.status_code = status
        return value
    monkeypatch.setattr(adapter.session, "get", get)
    scan, findings = run(adapter)
    assert scan.status == "COMPLETED"
    assert scan.execution_log[-1]["initial_response_coverage_limited"] is False
