"""Secrets must not survive through alternate finding fields or raw headers."""
import json

from docx import Document

from src.core.models import Finding, FindingConfidence, FindingSeverity, ScanRun
from src.reporting.docx_report import JSONExporter, ReportGenerator


def test_word_and_json_redact_every_finding_field_without_changing_source(tmp_path):
    scan = ScanRun(tool_name="Sentinel DAST", target="https://owned.test/?token=target-secret", raw_command="")
    item = Finding(title="Disclosure", source_tool="ZAP", target=scan.target,
                   description="Authorization: Basic BASIC_SECRET\nCookie: sid=COOKIE_ONE; arbitrary=COOKIE_TWO",
                   severity=FindingSeverity.HIGH, severity_source="Fixture",
                   confidence=FindingConfidence.MEDIUM, confidence_justification="Fixture",
                   evidence_raw=json.dumps({"request": "GET / HTTP/1.1\r\nAuthorization: Basic RAW_SECRET\r\nCookie: sid=RAW_COOKIE\r\n"}),
                   evidence_summary="password=SUMMARY_SECRET", remediation="api_key=FIX_SECRET")
    original = item.to_dict()
    json_path = JSONExporter(tmp_path).export_json(scan, [item], [])
    docx_path = ReportGenerator(tmp_path).generate_docx(scan, [item], [])
    document = Document(docx_path)
    word = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    for secret in ("target-secret", "BASIC_SECRET", "COOKIE_ONE", "COOKIE_TWO", "RAW_SECRET", "RAW_COOKIE", "SUMMARY_SECRET", "FIX_SECRET"):
        assert secret not in json_path.read_text(encoding="utf-8"), secret
        assert secret not in word, secret
    assert item.to_dict() == original


def test_structured_execution_log_redacts_headers_and_url_credentials(tmp_path, monkeypatch):
    from src.core import scan_logging
    monkeypatch.setattr(scan_logging, "LOGS_DIR", tmp_path)
    scan = ScanRun(tool_name="Fixture", target="https://name:URL_SECRET@owned.test/?api_key=QUERY_SECRET", raw_command="")
    scan.execution_log = [{"message": "Received Cookie: session=FIRST_COOKIE; other=SECOND_COOKIE\nAuthorization: Basic BASIC_LOG_SECRET",
                           "headers": {"X-ZAP-API-Key": "ZAP_API_SECRET"}}]
    path = scan_logging.write_execution_log(scan)
    text = path.read_text(encoding="utf-8")
    for secret in ("URL_SECRET", "QUERY_SECRET", "FIRST_COOKIE", "SECOND_COOKIE", "BASIC_LOG_SECRET", "ZAP_API_SECRET"):
        assert secret not in text
    assert json.loads(text)["scan"]["id"] == scan.id


def test_privacy_filter_preserves_non_secret_evidence():
    from src.core.privacy import sanitize_data
    data = {"description": "HTTP Basic authentication observed", "target": "https://owned.test/?id=12&token=PRIVATE", "request_count": 5}
    safe = sanitize_data(data)
    assert safe["description"] == data["description"]
    assert safe["request_count"] == 5
    assert "id=12" in safe["target"] and "PRIVATE" not in safe["target"]
