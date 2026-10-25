"""Labeled local corpus for discovery, policy evidence and false-positive resistance."""
import html
import json
import re
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit

import pytest

from src.adapters.javascript_analysis import analyze_flows, route_literals
from src.adapters.session_review import SessionReviewAdapter
from src.adapters.web_auditor import WebAuditorAdapter
from src.core.application_map import ApplicationMap
from src.core.models import ScanRun, FindingStatus
from src.core.privacy import sanitize_text


FLOW_CORPUS = [
    ('direct', 'el.innerHTML = location.hash;', True),
    ('alias', 'const q=location.search; const copy=q; el.innerHTML=copy;', True),
    ('minified', 'let a=location.hash;a=a.slice(1);document.write(a);', True),
    ('navigation', 'location.href=location.search;', True),
    ('quoted', 'const s="el.innerHTML=location.hash";', False),
    ('comment', '// el.innerHTML=location.hash\nconst safe=1;', False),
    ('block_comment', '/* el.innerHTML=location.hash; */', False),
    ('reassignment', 'let q=location.hash; q="safe"; el.innerHTML=q;', False),
    ('unrelated', 'const q=location.hash; el.innerHTML="safe";', False),
    ('safe_sink', 'el.textContent=location.hash;', False),
    ('sanitizer', 'el.innerHTML=DOMPurify.sanitize(location.hash);', False),
    ('scope', 'function first(){const q=location.hash;} function second(){el.innerHTML=q;}', False),
    ('regexp', 'const pattern=/location.hash document.write/;', False),
]


@pytest.mark.parametrize('name,code,expected', FLOW_CORPUS)
def test_labeled_static_corpus(name, code, expected):
    results = analyze_flows('https://owned.invalid/app.js', code, 'fixture')
    assert bool(results) is expected, name
    assert all(row.status == FindingStatus.POTENTIAL for row in results)
    assert all(row.confidence.value != 'CONFIRMED' for row in results)


@pytest.fixture
def target():
    observed = []
    class Target(BaseHTTPRequestHandler):
        def do_GET(self):
            parsed=urlsplit(self.path)
            observed.append(('GET', parsed.path, self.headers.get('Origin'), parse_qs(parsed.query)))
            kind='text/html'
            extra={}
            if parsed.path=='/app.js':
                kind='application/javascript'
                body='fetch("/api/profile?view=public"); fetch("/api/logout"); /* "/api/fake" */\n//# sourceMappingURL=maps/app.map'
            elif parsed.path=='/maps/app.map':
                kind='application/json'
                body=json.dumps({'version':3,'sources':['app.ts'],'mappings':'AAAA','sourcesContent':['const safe=1;']})
            elif parsed.path=='/fake.js':
                kind='application/javascript'
                body='//# sourceMappingURL=fake.map'
            elif parsed.path=='/fake.map':
                body='<html>"version":3,"sources":[]</html>'
            elif parsed.path in {'/reflect-cors','/static-cors','/failed-cors'}:
                body='public data'
                extra={'Access-Control-Allow-Origin':self.headers.get('Origin','') if parsed.path!='/static-cors' else 'https://sentinel.invalid', 'Access-Control-Allow-Credentials':'true'}
            elif parsed.path in {'/template','/literal','/fixed'}:
                value=parse_qs(parsed.query).get('q',[''])[0]
                body=value if parsed.path=='/literal' else 'SENTINEL-PROBE3431' if parsed.path=='/fixed' else re.sub(r'\{\{(\d+)\*(\d+)\}\}',lambda m:str(int(m[1])*int(m[2])),value)
            else:
                body='<script src="/app.js"></script><a href="/api/logout">Sign out</a>'
            encoded=body.encode()
            self.send_response(500 if parsed.path=='/failed-cors' else 200)
            self.send_header('Content-Type',kind)
            for key,value in extra.items():self.send_header(key,value)
            self.send_header('Content-Length',str(len(encoded)))
            self.end_headers();self.wfile.write(encoded)
        def do_POST(self):
            observed.append(('POST',self.path,None,{}));self.send_response(500);self.end_headers()
        def log_message(self,*args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Target)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:yield f'http://127.0.0.1:{server.server_port}',observed
    finally:server.shutdown();server.server_close();thread.join(3)


def test_live_js_map_and_route_provenance(target):
    base,observed=target
    adapter=WebAuditorAdapter()
    findings=adapter._audit_javascript_assets(base,[base+'/app.js?build=secret-value'])
    assert any(row.title=='Exposed JavaScript Source Map' for row in findings)
    assert adapter._discovered_endpoints==[base+'/api/logout',base+'/api/profile?view=public']
    snapshot=adapter._application_map.snapshot()
    assert any(row['relation']=='source_map_reference' for row in snapshot['relationships'])
    assert 'secret-value' not in json.dumps(snapshot)
    assert not any(row[1]=='/api/logout' for row in observed)
    assert not any(row[1]=='/api/fake' for row in observed)


def test_html_imitating_source_map_is_rejected(target):
    base,_=target
    findings=WebAuditorAdapter()._audit_javascript_assets(base,[base+'/fake.js'])
    assert not any(row.title=='Exposed JavaScript Source Map' for row in findings)


@pytest.mark.parametrize('path,expected',[('reflect-cors',True),('static-cors',False),('failed-cors',False)])
def test_live_cors_policy_corpus(target,path,expected):
    base,observed=target
    findings=WebAuditorAdapter()._check_cors(base+'/'+path)
    assert bool(findings) is expected
    assert all(row.status==FindingStatus.POTENTIAL for row in findings)
    assert len(observed)==3


@pytest.mark.parametrize('path,expected',[('template',True),('literal',False),('fixed',False)])
def test_live_arithmetic_corpus(target,path,expected):
    from urllib.parse import urlparse
    base,_=target
    adapter=WebAuditorAdapter()
    findings=adapter._verify_template_arithmetic(urlparse(base+'/'+path+'?q=original'),[('q','original')],'q','fixture',base+'/'+path)
    assert bool(findings) is expected
    if expected:
        proof=json.loads(findings[0].evidence_raw)
        assert len(proof['challenges'])==2
        assert proof['challenges'][0]['expected']!=proof['challenges'][1]['expected']


def test_discovery_graph_bounds_and_cross_origin():
    graph=ApplicationMap('https://owned.invalid')
    assert graph.add('https://elsewhere.invalid/private','api') is None
    assert graph.add('https://u:p@owned.invalid/private','api') is None
    for i in range(280):graph.add(f'https://owned.invalid/p{i}?token=not-for-storage','api','https://owned.invalid/app.js')
    value=graph.snapshot()
    assert len(value['nodes'])==250 and value['truncated']
    assert 'not-for-storage' not in json.dumps(value)


def test_bare_provider_tokens_and_unsigned_jwt_are_redacted():
    # Synthetic values only; never usable provider credentials.
    values=['ghp_'+'A'*36,'AKIA'+'Z'*16,'AIza'+'B'*35,'eyJhbGciOiJub25lIn0.eyJzdWIiOiJmaXh0dXJlIn0.']
    output=sanitize_text(' '.join(values))
    assert all(value not in output for value in values)


def test_static_corpus_metrics(tmp_path):
    tp=fp=tn=fn=0
    for name,code,expected in FLOW_CORPUS:
        actual=bool(analyze_flows('https://owned.invalid/app.js',code,'corpus'))
        tp+=int(actual and expected);fp+=int(actual and not expected)
        tn+=int(not actual and not expected);fn+=int(not actual and expected)
    record={'corpus':'13 labeled lexical JS cases; not a general vulnerability benchmark','tp':tp,'fp':fp,'tn':tn,'fn':fn,
            'precision':tp/(tp+fp) if tp+fp else 0,'recall':tp/(tp+fn) if tp+fn else 0}
    from pathlib import Path
    path=Path('cache/depth_static_metrics.json');path.parent.mkdir(exist_ok=True);path.write_text(json.dumps(record,indent=2))
    assert fp==fn==0


def test_new_findings_reach_ui_and_all_report_formats(target,tmp_path,monkeypatch):
    from PyQt6.QtWidgets import QApplication
    from src.ui.tabs import WebSecurityTab
    from src.core.database import DatabaseManager
    from src.reporting.export_service import export_assessment
    from docx import Document
    base,_=target
    findings=WebAuditorAdapter()._audit_javascript_assets(base,[base+'/app.js'])
    scan=ScanRun(tool_name='Local depth fixture',target=base,raw_command='',status='COMPLETED')
    db=DatabaseManager(tmp_path/'evidence.db');db.save_scan(scan)
    for row in findings:row.scan_id=scan.id;db.save_finding(row)
    app=QApplication.instance() or QApplication([])
    monkeypatch.setattr(DatabaseManager,'get_instance',classmethod(lambda cls:db))
    widget=WebSecurityTab()
    try:
        widget._on_scan_finished(scan,findings)
        assert widget.findings_table.rowCount()==len(findings)
        assert widget.scan_btn.text()=='Scan Website'
        assert widget.session_btn.isHidden()
        for fmt in ('json','markdown','html','docx'):
            path=export_assessment(db,scan.id,fmt,tmp_path/fmt)
            if fmt=='docx':
                document=Document(path)
                visible='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
            else:visible=path.read_text(encoding='utf-8')
            assert 'Source Map' in visible
    finally:widget.close();db.close()


def test_public_account_marker_cannot_validate_a_session(monkeypatch):
    from test_session_review import Target, target, plan
    # Exercise the existing local session fixture with a deliberately ambiguous canary.
    def public_canary(self):
        self.server.events.append((self.path,'synthetic'))
        body=b'ACCOUNT_ALICE ACCOUNT_BOB'
        self.send_response(200);self.send_header('Content-Length',str(len(body)));self.end_headers();self.wfile.write(body)
    monkeypatch.setattr(Target,'do_GET',public_canary)
    fixture=target.__wrapped__(monkeypatch)
    base,server=next(fixture)
    try:
        scan,findings=SessionReviewAdapter().run_scan(base,ScanRun(tool_name='Fixture',target=base,raw_command=''),plan())
        assert scan.status=='PARTIAL' and not findings
        assert all(path=='/account' for path,identity in server.events)
    finally:
        try:next(fixture)
        except StopIteration:pass


def test_transient_access_match_does_not_become_a_stable_finding(monkeypatch):
    from test_session_review import Target, target, plan
    original=Target.do_GET
    counts={}
    def flaky(self):
        cookie=self.headers.get('Cookie','')
        identity='alice' if 'QA_ALICE_SESSION' in cookie else 'bob' if 'QA_BOB_SESSION' in cookie else 'anonymous'
        if self.path=='/private' and identity!='alice':
            counts[identity]=counts.get(identity,0)+1
            self.server.exposed=counts[identity]==1
        return original(self)
    monkeypatch.setattr(Target,'do_GET',flaky)
    fixture=target.__wrapped__(monkeypatch);base,server=next(fixture)
    try:
        scan,findings=SessionReviewAdapter().run_scan(base,ScanRun(tool_name='Fixture',target=base,raw_command=''),plan())
        assert scan.status=='PARTIAL' and not findings
        assert counts=={'anonymous':2,'bob':2}
    finally:
        try:next(fixture)
        except StopIteration:pass


def test_automatic_probes_do_not_guess_an_adjacent_object(target,monkeypatch):
    base,observed=target
    adapter=WebAuditorAdapter()
    adapter._probe_query_parameters([{'url':base+'/item?id=1','names':['id']}],'fixture',None)
    assert not any(row[3].get('id')==['2'] for row in observed)
    assert all(row[0]=='GET' for row in observed)
    assert adapter._untested_checks
