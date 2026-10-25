import json
import os
import threading
from copy import deepcopy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from docx import Document

from src.adapters import session_review as module
from src.core.models import ScanRun
from src.reporting.docx_report import JSONExporter, ReportGenerator


def plan():
    return {"primary": {"mode": "cookie", "secret": "sid=QA_ALICE_SESSION", "marker": "ACCOUNT_ALICE"},
            "secondary": {"mode": "cookie", "secret": "sid=QA_BOB_SESSION", "marker": "ACCOUNT_BOB"},
            "canary_path": "/account", "read_only_confirmed": True,
            "cases": [{"path": "/private", "marker": "PRIVATE_ALICE_CONTENT", "anonymous_allowed": False, "secondary_allowed": False, "forbidden": ""}]}


class Target(BaseHTTPRequestHandler):
    def do_GET(self):
        cookie = self.headers.get("Cookie", "")
        auth = self.headers.get("Authorization", "")
        identity = "alice" if "QA_ALICE_SESSION" in cookie or auth == "Bearer QA_ALICE_TOKEN" else "bob" if "QA_BOB_SESSION" in cookie else "anonymous"
        self.server.events.append((self.path, identity))
        if self.path == "/account":
            status, body = (200, f"ACCOUNT_{identity.upper()}") if identity != "anonymous" and not self.server.expired else (401, "Please sign in")
            if self.server.login_page:
                status, body = 200, "Please sign in"
        elif self.path == "/public":
            status, body = 200, "PUBLIC_READ_ONLY_CONTENT"
        elif self.path == "/redirect":
            status, body = 302, "Moved"
        elif self.path == "/private":
            status, body = (200, "PRIVATE_ALICE_CONTENT " + self.server.extra) if identity == "alice" or self.server.exposed else (403, "Denied")
        else:
            status, body = 200, "<html>Loopback fixture</html>"
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        if self.path == "/redirect":
            self.send_header("Location", "http://127.0.0.1:1/never-follow")
        if self.path == "/public":
            self.send_header("Set-Cookie", "sid=QA_ALICE_SESSION; Path=/")
        content = body.encode()
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, *args):
        pass


@pytest.fixture
def target(monkeypatch):
    monkeypatch.setattr(module.time, "sleep", lambda seconds: None)
    server = ThreadingHTTPServer(("127.0.0.1", 0), Target)
    server.events, server.exposed, server.expired, server.extra = [], False, False, ""
    server.login_page = False
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        yield f"http://127.0.0.1:{server.server_port}/", server
    finally:
        server.shutdown()
        server.server_close()


def review(url, config):
    scan = ScanRun(tool_name=module.SessionReviewAdapter.name, target=url, raw_command="Read-only fixture")
    return module.SessionReviewAdapter().run_scan(url, scan, config)


def test_valid_sessions_and_denied_routes_have_observable_checks(target):
    url, server = target
    scan, findings = review(url, plan())
    assert scan.status == "COMPLETED" and not findings
    assert scan.request_count == len(server.events) == 6
    assert all(item["outcome"] == "PASS" for item in scan.execution_log[-1]["checks"])


def test_exposure_under_denied_identities_is_reported_without_secrets(target, tmp_path):
    url, server = target
    server.exposed = True
    scan, findings = review(url, plan())
    assert len(findings) == 2 and all(item.severity.value == "HIGH" for item in findings)
    assert all(item.status.value == "SCANNER_REPORTED" for item in findings)
    text = json.dumps(scan.to_dict()) + json.dumps([item.to_dict() for item in findings])
    for value in ("QA_ALICE_SESSION", "QA_BOB_SESSION", "ACCOUNT_ALICE", "ACCOUNT_BOB", "PRIVATE_ALICE_CONTENT"):
        assert value not in text
    root = ScanRun(tool_name="Fixture root", target=url, raw_command="", status="COMPLETED", execution_log=scan.execution_log)
    word = Document(ReportGenerator(tmp_path).generate_docx(root, findings, []))
    visible = "\n".join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
    assert "FAIL" in visible and "Session canary" in visible
    assert "QA_ALICE_SESSION" not in visible
    payload = JSONExporter(tmp_path).export_json(root, findings, []).read_text(encoding="utf-8")
    assert "QA_ALICE_SESSION" not in payload


def test_expired_session_is_inconclusive_and_paths_are_not_probed(target):
    url, server = target
    server.expired = True
    scan, findings = review(url, plan())
    assert scan.status == "PARTIAL" and not findings
    assert all(path == "/account" for path, identity in server.events)
    assert all(row["outcome"] == "INCONCLUSIVE" for row in scan.execution_log[-1]["checks"])


def test_bearer_primary_and_read_only_business_assertion(target):
    url, server = target
    config = plan()
    config["primary"].update(mode="bearer", secret="QA_ALICE_TOKEN")
    config["cases"][0]["forbidden"] = "LIMIT_EXCEEDED"
    server.extra = "LIMIT_EXCEEDED"
    scan, findings = review(url, config)
    assert scan.status == "COMPLETED" and len(findings) == 1
    assert findings[0].status.value == "POTENTIAL" and findings[0].severity.value == "LOW"


def test_redirect_and_body_limit_are_inconclusive(target, monkeypatch):
    url, server = target
    config = plan()
    config["cases"][0]["path"] = "/redirect"
    scan, findings = review(url, config)
    assert scan.status == "PARTIAL" and not findings
    assert not any("never-follow" in path for path, identity in server.events)
    monkeypatch.setattr(module, "MAX_BODY_BYTES", 5)
    scan, findings = review(url, plan())
    assert scan.status == "PARTIAL" and not findings


def test_anonymous_cookie_jar_cannot_turn_into_a_signed_in_role(target):
    url, server = target
    config = plan()
    config["cases"].insert(0, {"path": "/public", "marker": "PUBLIC_READ_ONLY_CONTENT", "anonymous_allowed": True, "secondary_allowed": True})
    scan, findings = review(url, config)
    assert scan.status == "COMPLETED" and not findings
    assert ("/private", "anonymous") in server.events


@pytest.mark.parametrize("mutation", ["origin", "fragment", "url_credentials", "downgrade", "short_marker", "same_session", "same_identity_marker", "newline", "read_only", "too_many"])
def test_invalid_plan_rejects_all_credentials_before_requests(monkeypatch, mutation):
    config = plan()
    if mutation == "origin": config["cases"][0]["path"] = "https://elsewhere.invalid/private"
    elif mutation == "fragment": config["cases"][0]["path"] = "/private#fragment"
    elif mutation == "url_credentials": config["cases"][0]["path"] = "https://name:secret@owned.invalid/private"
    elif mutation == "downgrade": config["cases"][0]["path"] = "http://owned.invalid/private"
    elif mutation == "short_marker": config["cases"][0]["marker"] = "x"
    elif mutation == "same_session": config["secondary"] = deepcopy(config["primary"])
    elif mutation == "same_identity_marker": config["secondary"]["marker"] = config["primary"]["marker"]
    elif mutation == "newline": config["primary"]["secret"] = "x\r\nAuthorization: x"
    elif mutation == "read_only": config["read_only_confirmed"] = False
    elif mutation == "too_many": config["cases"] *= 13
    monkeypatch.setattr(module.requests.Session, "get", lambda *a, **kw: pytest.fail("No request may be sent"))
    scan, findings = review("https://owned.invalid/", config)
    assert scan.status == "FAILED" and scan.request_count == 0 and not findings


def test_remote_http_credentials_are_rejected_before_network(monkeypatch):
    monkeypatch.setattr(module.requests.Session, "get", lambda *a, **kw: pytest.fail("No request may be sent"))
    scan, findings = review("http://owned.invalid/", plan())
    assert scan.status == "FAILED" and scan.request_count == 0


def test_worker_persists_child_review_and_exports_coverage_without_credentials(target, tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QApplication
    from src.ui.tabs import WebAuditWorker
    from src.core.database import DatabaseManager
    from src.core.models import ToolHealthInfo
    from src.reporting.export_service import export_assessment
    app = QApplication.instance() or QApplication([])
    url, server = target
    server.exposed = True
    database = DatabaseManager(tmp_path / "integrated.db")
    monkeypatch.setattr(DatabaseManager, "get_instance", classmethod(lambda cls: database))
    monkeypatch.setattr("src.ui.tabs.ZapAdapter.detect", lambda self: ToolHealthInfo(name="ZAP", installed=False))
    monkeypatch.setattr("src.ui.tabs.NmapAdapter.detect", lambda self: ToolHealthInfo(name="Nmap", installed=False))
    monkeypatch.setattr("src.ui.tabs.write_execution_log", lambda scan: tmp_path / "log.json")
    worker = WebAuditWorker(url, session_plan=plan())
    results = []
    worker.finished_signal.connect(lambda scan, findings: results.append((scan, findings)))
    worker.run()
    assert results[0][0].status == "COMPLETED"
    assert worker.session_plan is None
    children = [scan for scan in database.list_scans() if scan.parent_scan_id == results[0][0].id]
    assert len(children) == 1 and children[0].request_count == 9
    report = export_assessment(database, results[0][0].id, "json", tmp_path / "exports")
    payload = json.loads(report.read_text(encoding="utf-8"))
    assert any(row["operation"] == "Session/access review" for row in payload["scan"]["execution_log"])
    assert len([finding for finding in payload["findings"] if finding["source_tool"] == module.SessionReviewAdapter.name]) == 2
    assert "QA_ALICE_SESSION" not in report.read_text(encoding="utf-8")
    assert "QA_BOB_SESSION" not in report.read_text(encoding="utf-8")
    database.close()


def test_setup_dialog_is_validated_and_can_be_cleared(target, monkeypatch):
    from PyQt6.QtWidgets import QApplication, QDialog
    from src.ui.session_review_dialog import SessionReviewDialog
    app = QApplication.instance() or QApplication([])
    url, server = target
    dialog = SessionReviewDialog(url, plan())
    dialog.apply_setup()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert dialog.plan["primary"]["secret"] == "sid=QA_ALICE_SESSION"
    dialog.clear_setup()
    assert dialog.plan is None and not dialog.profiles["primary"][1].text()
    dialog.close()


def test_http_200_login_page_does_not_validate_a_session(target):
    url, server = target
    server.login_page = True
    scan, findings = review(url, plan())
    assert scan.status == "PARTIAL" and not findings and scan.response_status == 200
    assert all(path == "/account" for path, identity in server.events)


def test_primary_content_marker_mismatch_is_not_a_pass(target):
    url, server = target
    config = plan()
    config["cases"][0]["path"] = "/public"
    scan, findings = review(url, config)
    assert scan.status == "PARTIAL" and not findings
    assert not any(path == "/public" and identity != "alice" for path, identity in server.events)


def test_transport_error_does_not_leak_credential_text(monkeypatch):
    import requests
    monkeypatch.setattr(module.time, "sleep", lambda seconds: None)
    def fail(*args, **kwargs):
        raise requests.Timeout("Cookie: QA_TRANSPORT_SECRET")
    monkeypatch.setattr(module.requests.Session, "get", fail)
    scan, findings = review("https://owned.invalid/", plan())
    assert scan.status == "PARTIAL" and not findings
    assert "QA_TRANSPORT_SECRET" not in json.dumps(scan.to_dict())
    assert all(item["error_type"] == "Timeout" for item in scan.execution_log[-1]["checks"])


def test_session_diagnostic_is_loopback_only_and_persists_local_findings(target, tmp_path):
    from main import _verify_local_session_review
    url, server = target
    output = tmp_path / "verification.json"
    config = tmp_path / "plan.json"
    config.write_text(json.dumps(plan()))
    assert _verify_local_session_review("https://owned.invalid", str(config), str(output)) == 1
    assert not server.events
    server.exposed = True
    assert _verify_local_session_review(url, str(config), str(output)) == 0
    payload = json.loads(output.read_text())
    assert len(payload["findings"]) == 2
    assert payload["persistence_verified"] and payload["export_verified"]
    assert "QA_ALICE_SESSION" not in output.read_text()
