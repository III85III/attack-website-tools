"""Local fixture imports never launch scanners or claim independent validation."""
import hashlib
import json

import pytest

from src.adapters import report_import_adapter as module
from src.core.models import FindingSeverity, FindingStatus
from src.reporting.docx_report import JSONExporter, ReportGenerator
from docx import Document


def event(**changes):
    value = {"template-id": "fixture-CVE-2026-12345", "type": "http",
             "info": {"name": "Fixture exposure", "severity": "critical",
                      "description": "Local import fixture", "remediation": "Remove exposure",
                      "classification": {"cwe-id": ["CWE-200"]}},
             "matched-at": "https://owned.test/path?token=QUERY_PRIVATE",
             "matcher-status": True,
             "request": "GET / HTTP/1.1\r\nAuthorization: Basic REQUEST_PRIVATE\r\n",
             "response": "HTTP/1.1 200 OK\r\nSet-Cookie: sid=COOKIE_PRIVATE; other=SECOND_PRIVATE\r\n",
             "template-encoded": "NEVER_RETAIN_TEMPLATE", "curl-command": "NEVER_EXECUTE_COMMAND"}
    value.update(changes)
    return value


def write_report(tmp_path, payload, jsonl=False):
    path = tmp_path / ("fixture.jsonl" if jsonl else "fixture.json")
    path.write_text("\n".join(json.dumps(row) for row in payload) if jsonl else json.dumps(payload), encoding="utf-8")
    return path


@pytest.mark.parametrize("shape", ["object", "array", "jsonl"])
def test_nuclei_formats_provenance_redaction_and_word_json(tmp_path, shape, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Report import must not make requests or start processes")
    monkeypatch.setattr("requests.sessions.Session.request", forbidden)
    monkeypatch.setattr("subprocess.Popen", forbidden)
    payload = event() if shape == "object" else [event(), event()]
    path = write_report(tmp_path, payload, shape == "jsonl")
    scan, findings = module.ReportImportAdapter().import_file(path)
    assert scan.status == "COMPLETED" and scan.request_count == 0
    assert len(findings) == scan.findings_count == 1
    finding = findings[0]
    assert finding.severity == FindingSeverity.CRITICAL
    assert finding.status == FindingStatus.SCANNER_REPORTED
    assert finding.cwe_id == "CWE-200" and finding.cve_id == "CVE-2026-12345"
    assert finding.tool_version == "Not supplied by report"
    assert "not independently verified" in scan.completion_reason
    receipt = scan.execution_log[-1]
    assert receipt["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert receipt["source_bytes"] == path.stat().st_size
    assert receipt["network_requests"] == 0
    json_path = JSONExporter(tmp_path).export_json(scan, findings, [])
    word_path = ReportGenerator(tmp_path).generate_docx(scan, findings, [])
    document = Document(word_path)
    word = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    combined = json_path.read_text(encoding="utf-8") + word + json.dumps(finding.to_dict())
    for secret in ("QUERY_PRIVATE", "REQUEST_PRIVATE", "COOKIE_PRIVATE", "SECOND_PRIVATE", "NEVER_RETAIN_TEMPLATE", "NEVER_EXECUTE_COMMAND"):
        assert secret not in combined
    assert "Fixture exposure" in word


def test_failed_matchers_and_errors_do_not_become_vulnerabilities(tmp_path):
    path = write_report(tmp_path, [event(**{"matcher-status": False}), event(error="request failed")], True)
    scan, findings = module.ReportImportAdapter().import_file(path)
    assert not findings and scan.status == "PARTIAL"
    assert scan.execution_log[0]["failed_or_error_records_skipped"] == 2


@pytest.mark.parametrize("payload", [[], [event(), {}], event(info=[]), event(**{"template-id": ""}),
                                      event(**{"matcher-status": "false"}), event(**{"matched-at": ""})])
def test_invalid_nuclei_records_reject_whole_import(tmp_path, payload):
    with pytest.raises(ValueError):
        module.ReportImportAdapter().import_file(write_report(tmp_path, payload))


def test_malformed_jsonl_rejects_line_without_partial_success(tmp_path):
    path = write_report(tmp_path, [event()], True)
    path.write_text(path.read_text() + "\n{broken", encoding="utf-8")
    with pytest.raises(ValueError, match="line 2"):
        module.ReportImportAdapter().import_file(path)


@pytest.mark.parametrize("kind", ["nuclei", "jsonl", "generic", "sarif", "hydra"])
def test_excess_records_are_rejected_not_silently_truncated(tmp_path, monkeypatch, kind):
    monkeypatch.setattr(module, "MAX_IMPORTED_FINDINGS", 1)
    payload = {
        "nuclei": [event(), event()], "jsonl": [event(), event()],
        "generic": {"findings": [{"title": "One"}, {"title": "Two"}]},
        "sarif": {"version": "2.1.0", "runs": [{"results": [{"ruleId": "one"}, {"ruleId": "two"}]}]},
        "hydra": {"generator": {"software": "hydra"}, "results": [{"host": "one"}, {"host": "two"}]},
    }[kind]
    with pytest.raises(ValueError, match="record import limit"):
        module.ReportImportAdapter().import_file(write_report(tmp_path, payload, kind == "jsonl"))


def test_explicit_empty_generic_report_has_honest_import_receipt(tmp_path):
    scan, findings = module.ReportImportAdapter().import_file(write_report(tmp_path, {"findings": []}))
    assert not findings and scan.status == "COMPLETED"
    assert scan.execution_log[-1]["imported_findings"] == 0
    assert "source scan" in scan.completion_reason
