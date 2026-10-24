"""Generate clear Persian security reports from saved Sentinel evidence."""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from dataclasses import replace

from src.core.config import REPORTS_DIR
from src.core.privacy import sanitize_data, sanitize_text
from src.core.models import Finding, ScanRun, ToolHealthInfo


_FONT = "Tahoma"
_NAVY = "17365D"
_BLUE = "D9EAF7"
_PALE = "F3F6FA"
_BORDER = "D9D9D9"
_SEVERITY_ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "INFO": 4}
_SEVERITY_FA = {
    "CRITICAL": "بحرانی",
    "HIGH": "زیاد",
    "MEDIUM": "متوسط",
    "LOW": "کم",
    "INFO": "اطلاعاتی",
}
_STATUS_FA = {
    "VALIDATED": "اعتبارسنجی‌شده",
    "SCANNER_REPORTED": "گزارش‌شده توسط اسکنر؛ نیازمند بررسی مستقل",
    "OBSERVED": "مشاهده‌شده؛ الزاماً آسیب‌پذیری نیست",
    "POTENTIAL": "احتمالی؛ نیازمند بررسی",
    "FALSE_POSITIVE": "مثبت کاذب اعلام‌شده",
    "NOT_VERIFIED": "بررسی‌نشده",
    "FAILED_COLLECTION": "جمع‌آوری شواهد ناموفق",
}
_CONFIDENCE_FA = {
    "CONFIRMED": "تأییدشده",
    "HIGH": "زیاد",
    "MEDIUM": "متوسط",
    "LOW": "کم",
    "TENTATIVE": "مقدماتی",
}
_CONFIDENCE_MEANING = {
    "CONFIRMED": "منبع برای این وضعیت شاهد مستقیم ثبت کرده است؛ این عنوان الزاماً به معنی اثبات اثر کامل حمله نیست.",
    "HIGH": "شواهد قوی است، اما دامنهٔ اثر را با توجه به برنامهٔ واقعی بررسی کنید.",
    "MEDIUM": "نشانهٔ قابل‌پیگیری وجود دارد و بررسی دستی می‌تواند نتیجه را دقیق‌تر کند.",
    "LOW": "این مورد یک سرنخ اولیه است و پیش از تصمیم اصلاحی نیاز به بازبینی دارد.",
    "TENTATIVE": "این مورد مقدماتی است؛ تا بررسی بیشتر آن را قطعی در نظر نگیرید.",
}
_SCAN_STATUS_FA = {
    "RUNNING": "در حال اجرا",
    "COMPLETED": "تکمیل‌شده",
    "PARTIAL": "ناقص",
    "FAILED": "ناموفق",
    "CANCELLED": "لغوشده",
}
_SEVERITY_MEANING = {
    "CRITICAL": "اولویت رسیدگی فوری. پیش از بستن این گزارش، شواهد و اثر عملیاتی این مورد را بررسی کنید.",
    "HIGH": "اولویت بالا. در نزدیک‌ترین بازهٔ اصلاح، مورد را بازبینی و رفع کنید.",
    "MEDIUM": "در برنامهٔ اصلاح قرار دهید و اثر آن را با توجه به سامانهٔ خود بسنجید.",
    "LOW": "بهبود امنیتی پیشنهادی؛ در فرصت مناسب پیکربندی را سخت‌سازی کنید.",
    "INFO": "این مورد به‌تنهایی آسیب‌پذیری نیست؛ برای شناخت سطح حمله یا زمینهٔ بررسی ثبت شده است.",
}
_SENSITIVE_KEY = re.compile(r"password|passwd|secret|token|api[_-]?key|authorization|cookie", re.I)
_SECRET_TEXT = re.compile(
    r'''(?i)(["']?\b(?:password|passwd|secret|token|api[_-]?key|authorization|cookie)["']?\s*[:=]\s*["']?)([^"'\s,;}]+)'''
)


def _value(value: Any) -> str:
    return value.value if hasattr(value, "value") else str(value or "")


def _redact_string(value: str) -> str:
    value = _SECRET_TEXT.sub(r"\1[REDACTED]", value)
    value = re.sub(r"(?i)(set-cookie\s*:\s*)[^\r\n]+", r"\1[REDACTED]", value)
    value = re.sub(r"(?i)\bBearer\s+[A-Za-z0-9._~+/=-]+", "Bearer [REDACTED]", value)
    value = re.sub(r"\beyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\b", "[REDACTED_JWT]", value)
    return value


def _safe_object(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            str(key): "[REDACTED]" if _SENSITIVE_KEY.search(str(key)) else _safe_object(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_safe_object(item) for item in value[:100]]
    if isinstance(value, str):
        return _redact_string(value[:2500])
    return value


def _safe_evidence(finding: Finding) -> str:
    raw = str(finding.evidence_raw or "")[:12000]
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return _redact_string(raw[:2500]) or "شاهد خامی برای این مورد ثبت نشده است."
    return json.dumps(_safe_object(parsed), ensure_ascii=False, indent=2)[:3500]


def _plain_finding_text(finding: Finding) -> tuple[str, str, str]:
    """Translate Sentinel's common checks into short Persian reader guidance."""
    title = finding.title or "یافتهٔ امنیتی"
    lower = title.lower()
    if lower == "api authentication declaration differs from anonymous response":
        return ("ناسازگاری احتمالی سیاست احراز هویت API", "تعریف API برای این عملیات احراز هویت می‌خواهد، اما درخواست بدون اعتبارنامه و بدون کوکی قبلی پاسخ موفق و غیرخالی دریافت کرده است. این نتیجه به‌تنهایی دسترسی به دادهٔ خصوصی را ثابت نمی‌کند؛ ممکن است مستندات قدیمی یا پاسخ عمداً عمومی باشد.", "سیاست واقعی و محتوای حفاظت‌شده را با مالک سامانه و حساب آزمایشی معتبر بررسی کنید. اگر داده خصوصی است، احراز هویت را پیش از پاسخ اعمال کنید؛ اگر مسیر عمومی است، تعریف API را اصلاح کنید.")
    if lower == "browser-confirmed reflected xss":
        return ("اجرای XSS بازتابی در مرورگر تأیید شد", "ورودی پارامتر URL توانست کد JavaScript بی‌خطر را در صفحهٔ سایت اجرا کند. لینک ساخته‌شده توسط مهاجم ممکن است کد را در مبدأ سایت اجرا کند؛ دسترسی به حساب یا سرقت داده در این بررسی آزمایش نشده است.", "در محل نمایش داده، خروجی را متناسب با بافت HTML یا ویژگی آن encode کنید؛ از APIهای امن DOM مانند textContent استفاده کنید. CSP دفاع تکمیلی است. پس از اصلاح، همان پارامتر و مسیر را دوباره بررسی کنید.")
    if lower == "potential unescaped reflected input":
        return ("بازتاب ورودی بدون escape؛ نامزد XSS", "نشانگر ورودی در HTML دیده شد، اما اجرای JavaScript هنوز اثبات نشده است. نتیجهٔ تلاش مرورگری و محدودیت آن در شواهد آمده است.", "بافت نمایش داده را بررسی و خروجی را متناسب با آن encode کنید؛ برای نتیجه قطعی، اجرای واقعی باید تأیید شود.")
    if "insecure unencrypted transport" in lower:
        return ("ارتباط وب‌سایت رمزگذاری نشده است", "اطلاعاتی که کاربر به این سایت می‌فرستد ممکن است در مسیر شبکه خوانده یا تغییر داده شود. این نتیجه از استفادهٔ HTTP به‌دست آمده و به معنی کشف سرقت اطلاعات نیست.", "برای سایت HTTPS فعال کنید و پس از اطمینان از کارکرد آن، کاربران را به نسخهٔ امن هدایت کنید.")
    if "password form served over unencrypted http" in lower:
        return ("فرم گذرواژه روی ارتباط ناامن نمایش داده شده است", "صفحه‌ای که ورودی گذرواژه دارد با HTTP باز شده است. Sentinel فرم را ارسال نکرده و هیچ گذرواژه‌ای آزمایش یا ذخیره نشده است.", "صفحهٔ ورود و مقصد ارسال فرم را فقط از HTTPS ارائه کنید.")
    if "content-security-policy" in lower:
        return ("سیاست امنیت محتوای مرورگر تنظیم نشده است", "مرورگر برای محدودکردن منبع اسکریپت‌ها و محتوای صفحه، سیاست CSP دریافت نکرده است. این مورد به‌تنهایی وجود XSS را ثابت نمی‌کند، اما یک لایهٔ دفاعی مهم غایب است.", "یک CSP متناسب با منابع واقعی سایت طراحی و ابتدا در حالت گزارش‌گیری آزمایش کنید؛ سپس آن را اعمال کنید.")
    if "strict-transport-security" in lower or " hsts" in lower:
        return ("سیاست HSTS دریافت نشد", "مرورگرها از این پاسخ دستور ماندگاری روی HTTPS دریافت نکردند. کاربرانی که ابتدا با HTTP به سایت می‌رسند ممکن است در معرض کاهش امنیت اتصال باشند.", "پس از فعال‌سازی کامل HTTPS، هدر Strict-Transport-Security را با مدت‌زمان مناسب تنظیم کنید.")
    if "clickjacking" in lower or "x-frame-options" in lower:
        return ("محدودیت بارگذاری سایت داخل قاب دریافت نشد", "پاسخ سایت مانع نمایش صفحه در قاب یک سایت دیگر نشد. در صفحات تعاملی، این وضعیت می‌تواند زمینهٔ فریب کلیک کاربر را فراهم کند؛ آسیب‌پذیری قطعی از این هدر به‌تنهایی نتیجه نمی‌شود.", "برای صفحات مناسب، X-Frame-Options یا دستور frame-ancestors در CSP تنظیم کنید.")
    if "x-content-type-options" in lower:
        return ("جلوگیری از حدس نوع محتوا تنظیم نشده است", "مرورگر ممکن است نوع بعضی پاسخ‌ها را حدس بزند. این تنظیم به‌ویژه برای فایل‌هایی که از کاربر دریافت می‌شوند اهمیت دارد.", "هدر X-Content-Type-Options را با مقدار nosniff ارسال کنید.")
    if "detailed server banner" in lower:
        return ("نسخه یا جزئیات وب‌سرور در پاسخ دیده می‌شود", "پاسخ وب‌سایت اطلاعاتی دربارهٔ نرم‌افزار سرور نشان داده است. این افشا به‌تنهایی آسیب‌پذیری را ثابت نمی‌کند، اما شناسایی فناوری را برای مهاجم آسان‌تر می‌کند.", "نمایش نسخه و جزئیات غیرضروری را در تنظیمات وب‌سرور کاهش دهید و وصله‌ها را به‌روز نگه دارید.")
    if "x-powered-by" in lower:
        return ("فناوری سمت سرور در پاسخ اعلام شده است", "هدر پاسخ نام یا نسخهٔ فناوری پشت وب‌سایت را افشا می‌کند. این مورد به‌تنهایی ضعف قابل‌استفاده را اثبات نمی‌کند.", "هدر X-Powered-By را در تنظیمات برنامه یا وب‌سرور حذف کنید.")
    if "missing httponly flag" in lower:
        cookie = re.search(r"Cookie '([^']+)'", title)
        return ("کوکی بدون ویژگی HttpOnly تنظیم شده است", f"کد JavaScript صفحه می‌تواند کوکی «{cookie.group(1) if cookie else 'موردنظر'}» را بخواند. اگر ضعف دیگری مانند XSS وجود داشته باشد، خطر افشای نشست افزایش می‌یابد.", "برای کوکی نشست و داده‌های حساس، ویژگی HttpOnly را فعال کنید.")
    if "missing secure flag" in lower:
        cookie = re.search(r"Cookie '([^']+)'", title)
        return ("کوکی HTTPS ویژگی Secure ندارد", f"کوکی «{cookie.group(1) if cookie else 'موردنظر'}» بدون پرچم Secure ارسال شده است. اگر نسخهٔ HTTP سایت در دسترس باشد، مرورگر ممکن است آن را روی اتصال ناامن نیز بفرستد.", "برای کوکی‌هایی که باید فقط روی HTTPS ارسال شوند، ویژگی Secure را فعال کنید.")
    if "git metadata" in lower or "public .git repository" in lower:
        return ("نشانه‌ای از دسترسی عمومی به فایل Git دیده شد", "پاسخ سرور با الگوی فایل HEAD مخزن Git مطابقت داشت. این اسکن فقط همین نشانه را بررسی کرده و دانلود کامل کد یا تاریخچهٔ مخزن را تأیید نکرده است.", "دسترسی وب به پوشهٔ .git را مسدود کنید و بررسی کنید اطلاعات حساس در تاریخچهٔ مخزن افشا نشده باشد.")
    if "environment configuration" in lower or "public .env configuration" in lower:
        return ("فایل پیکربندی محیطی در دسترس وب دیده شد", "پاسخ، الگوی متغیرهای محیطی را داشت؛ محتوای فایل و مقادیر محرمانه در گزارش ذخیره نشده‌اند.", "دسترسی وب به فایل .env را مسدود کنید، آن را بیرون از ریشهٔ وب نگه دارید و اگر احتمال افشا هست، کلیدها را تعویض کنید.")
    if "robots.txt discloses" in lower:
        return ("فایل robots.txt مسیرهایی را نام برده است", "این فایل مسیرهای ذکرشده را پنهان نمی‌کند و کنترل دسترسی نیست. هر مسیر حساس باید جداگانه با مجوز سمت سرور محافظت شود.", "دسترسی به مسیرهای حساس را با احراز هویت و مجوزدهی سمت سرور کنترل کنید.")
    if "web attack surface mapped" in lower:
        return ("مسیرهای قابل‌دسترسی وب‌سایت فهرست شدند", "این یک مشاهدهٔ اطلاعاتی است، نه آسیب‌پذیری. مسیرهای فهرست‌شده را مرور کنید و مطمئن شوید صفحه‌های حساس مجوزدهی سمت سرور دارند.", "مسیرهای فهرست‌شده را با محدودهٔ مورد انتظار برنامه مقایسه کنید.")
    if "open port" in lower:
        return ("یک پورت شبکه پاسخ می‌دهد", "روی این میزبان سرویسی در دسترس است. بازبودن پورت به‌تنهایی آسیب‌پذیری نیست؛ نیاز کسب‌وکار، نسخهٔ سرویس و محدودیت دسترسی را بررسی کنید.", "اگر سرویس لازم نیست آن را ببندید؛ در غیر این صورت دسترسی را به مبدأهای مجاز محدود و سرویس را به‌روز کنید.")
    if "connection failure" in lower or "target rejected" in lower:
        return ("اسکن نتوانست شواهد کافی جمع‌آوری کند", "این مورد آسیب‌پذیری نیست؛ اتصال یا ورودی هدف ناموفق بود و نتیجهٔ امنیتی کاملی دربارهٔ سایت نمی‌دهد.", "URL، دسترس‌پذیری سرویس و خطای ثبت‌شده را بررسی کنید و بعد اسکن را تکرار کنید.")
    return (title, finding.description or "توضیحی در گزارش منبع ثبت نشده است.", finding.remediation or "راهکار مشخصی در گزارش منبع ثبت نشده است.")


def _local_time(value: Optional[str]) -> str:
    if not value:
        return "ثبت نشده"
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone().strftime("%Y-%m-%d  %H:%M") + "  (ساعت محلی)"
    except (ValueError, TypeError):
        return str(value)[:80]


def _set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12


def _style_run(run, size: int = 10, bold: bool = False, color: str = "222222") -> None:
    run.font.name = _FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), _FONT)
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), _BORDER)


def _prepare_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    tbl_pr.append(bidi)
    _set_table_borders(table)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _set_cell_text(cell, text: Any, *, bold: bool = False, color: str = "222222", size: int = 9, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_rtl(paragraph, align)
    run = paragraph.add_run(str(text or ""))
    _style_run(run, size=size, bold=bold, color=color)


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _add_table(doc, headers: List[str], rows: List[List[str]], widths: List[float], *, font_size: int = 9):
    table = doc.add_table(rows=1, cols=len(headers))
    _prepare_table(table)
    for index, title in enumerate(headers):
        table.columns[index].width = Inches(widths[index])
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, _NAVY)
        _set_cell_text(cell, title, bold=True, color="FFFFFF", size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.width = Inches(widths[index])
    _repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[index])
            cells[index].width = Inches(widths[index])
            if row_index % 2:
                _set_cell_shading(cells[index], _PALE)
            _set_cell_text(cells[index], value, size=font_size)
    return table


def _heading(doc, text: str, level: int):
    paragraph = doc.add_heading(text, level=level)
    _set_rtl(paragraph)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, value, end):
        run = paragraph.add_run()
        run._r.append(element)
        _style_run(run, size=8, color="667085")


def _configure_document(doc) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = _FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    _set_rtl(header, WD_ALIGN_PARAGRAPH.RIGHT)
    _style_run(header.add_run("III85III QOLAM SENTINEL   |   گزارش امنیتی"), size=8, color="667085")
    footer = section.footer.paragraphs[0]
    _set_rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)
    _style_run(footer.add_run("محرمانه   •   صفحه "), size=8, color="667085")
    _add_field(footer, "PAGE")


class ReportGenerator:
    """Create a Persian report that separates observations from validated risks."""

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or REPORTS_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_docx(
        self,
        scan_run: Optional[ScanRun],
        findings: List[Finding],
        tools_health: List[ToolHealthInfo],
        target_name: str = "",
        report_title: str = "گزارش ارزیابی امنیتی",
    ) -> Path:
        # Redact every visible field, not only the raw-evidence table. Work on
        # copies so exported artifacts cannot alter the stored assessment.
        scan_run = ScanRun.from_dict(sanitize_data(scan_run.to_dict())) if scan_run else None
        findings = [Finding.from_dict(sanitize_data(finding.to_dict())) for finding in findings]
        tools_health = [replace(tool, name=sanitize_text(tool.name),
                                version=sanitize_text(tool.version) if tool.version else None,
                                path=sanitize_text(tool.path) if tool.path else None,
                                status_message=sanitize_text(tool.status_message)) for tool in tools_health]
        target_name, report_title = sanitize_text(target_name), sanitize_text(report_title)
        doc = docx.Document()
        _configure_document(doc)
        target = target_name or (scan_run.target if scan_run else "هدف مشخص نشده")
        generated_at = datetime.now().astimezone().strftime("%Y-%m-%d  %H:%M") + "  (ساعت محلی)"

        title = doc.add_paragraph(style="Title")
        _set_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
        _style_run(title.add_run(report_title or "گزارش ارزیابی امنیتی"), size=22, bold=True, color="000000")
        subtitle = doc.add_paragraph()
        _set_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
        _style_run(subtitle.add_run("گزارش مستند یافته‌ها و اولویت‌های اصلاح"), size=11, color="44546A")
        generated = doc.add_paragraph()
        _set_rtl(generated, WD_ALIGN_PARAGRAPH.CENTER)
        _style_run(generated.add_run(f"تاریخ تهیه: {generated_at}"), size=9, color="667085")

        _heading(doc, "خلاصهٔ مدیریتی", level=1)
        counts = {severity: 0 for severity in _SEVERITY_ORDER}
        false_positive_count = 0
        for finding in findings:
            if _value(finding.status).upper() == "FALSE_POSITIVE":
                false_positive_count += 1
                continue
            severity = _value(finding.severity).upper()
            if severity in counts:
                counts[severity] += 1
        total_risks = sum(counts[key] for key in ("CRITICAL", "HIGH", "MEDIUM", "LOW"))
        if counts["CRITICAL"]:
            summary = f"در این ارزیابی {counts['CRITICAL']} مورد بحرانی ثبت شده است. رسیدگی را با بررسی فوری شواهد و کاهش خطر این موارد آغاز کنید."
        elif counts["HIGH"]:
            summary = f"در این ارزیابی مورد بحرانی ثبت نشده، اما {counts['HIGH']} مورد با اولویت زیاد وجود دارد و باید در برنامهٔ اصلاح نزدیک قرار گیرد."
        elif total_risks:
            summary = f"در این ارزیابی مورد بحرانی یا زیاد ثبت نشده است. {total_risks} مورد با شدت متوسط یا کم برای بازبینی و بهبود پیکربندی ثبت شده است."
        else:
            summary = "در داده‌های این ارزیابی، موردی با شدت بحرانی تا کم ثبت نشده است. این نتیجه به معنی اثبات امنیت کامل سامانه نیست؛ دامنه و محدودیت آزمون را در ادامه ببینید."
        if false_positive_count:
            summary += f" تعداد {false_positive_count} موردی که به‌عنوان مثبت کاذب علامت‌گذاری شده‌اند، از شمارش ریسک فعال کنار گذاشته شدند و همچنان در فهرست یافته‌ها دیده می‌شوند."
        self._paragraph(doc, summary)

        verification_counts = {}
        for finding in findings:
            state = _value(finding.status).upper()
            verification_counts[state] = verification_counts.get(state, 0) + 1
        if verification_counts:
            _heading(doc, "تفکیک وضعیت شواهد", level=2)
            _add_table(doc, ["وضعیت ثبت‌شده", "تعداد"],
                       [[_STATUS_FA.get(state, state), str(count)] for state, count in sorted(verification_counts.items())],
                       [5.6, 1.2], font_size=9)
            self._paragraph(doc, "شدت با اعتبارسنجی یکسان نیست. وضعیت اعتبارسنجی‌شده ممکن است توسط ارزیاب ثبت شده باشد؛ برای اثبات روش تأیید، شواهد هر مورد را بازبینی کنید.")

        meta_rows = [
            ["هدف", target],
            ["اسکن", scan_run.tool_name if scan_run else "به اسکن مشخصی پیوند ندارد"],
            ["شروع", _local_time(scan_run.start_time) if scan_run else "ثبت نشده"],
            ["پایان و وضعیت", f"{_local_time(scan_run.end_time) if scan_run else 'ثبت نشده'}  |  {_SCAN_STATUS_FA.get(_value(scan_run.status).upper(), _value(scan_run.status)) if scan_run else 'نامشخص'}"],
            ["شناسهٔ اسکن", scan_run.id if scan_run else "ثبت نشده"],
        ]
        if scan_run:
            meta_rows.extend([
                ["درخواست‌های HTTP بررسی اولیه", str(scan_run.request_count)],
                ["وضعیت پاسخ نخست", str(scan_run.response_status or "ثبت نشده")],
                ["مدت اجرا (میلی‌ثانیه)", str(scan_run.duration_ms)],
                ["علت وضعیت اسکن", scan_run.completion_reason or scan_run.error_message or "ثبت نشده"],
                ["مسیر لاگ ساختاریافته", next((item.get("path") for item in reversed(scan_run.execution_log) if item.get("path")), "در JSON / پایگاه داده")],
            ])
        _add_table(doc, ["مشخصات", "مقدار"], meta_rows, [1.45, 5.35], font_size=9)

        if scan_run:
            api_review = next((entry for entry in scan_run.execution_log if entry.get("operation") == "API coverage"), None)
            if api_review:
                _heading(doc, "پوشش مسیرهای API", level=2)
                self._paragraph(doc, f"تعریف API: {api_review.get('spec_version', 'نامشخص')} | عملیات اعلام‌شده: {api_review.get('declared_operations', 'نامشخص')} | وضعیت: {api_review.get('state', 'نامشخص')}")
                self._paragraph(doc, f"درخواست HTTP این بخش: {api_review.get('request_count', 0)} | درخواست مرورگر این بخش: {api_review.get('browser_request_count', 0)}؛ هر دو جدا از شمارش HTTP صفحهٔ اصلی‌اند.")
                self._paragraph(doc, "SHA-256 تعریف ورودی: " + str(api_review.get("source_sha256", "ثبت نشده")))
                rows = [[str(item.get('method', '')), str(item.get('path', '')), str(item.get('http_status', '')),
                         str(item.get('outcome', '')), str(item.get('reason', ''))] for item in api_review.get('checks', [])]
                if rows:
                    _add_table(doc, ["روش", "مسیر", "HTTP", "نتیجه", "علت / محدودیت"], rows, [.55, 2.15, .6, 1.0, 2.5], font_size=8)
                self._paragraph(doc, "COLLECTED فقط دریافت موفق پاسخ است؛ به معنی نبود آسیب‌پذیری یا تأیید کنترل دسترسی نیست. NOT_TESTED و INCONCLUSIVE پوشش امنیتی تأییدشده ندارند. منابع دارای شناسه، POST/JSON، نقش‌ها و منطق کسب‌وکار نیازمند سناریوی صریح‌اند.")
            browser_review = next((entry for entry in scan_run.execution_log if entry.get("operation") == "Browser XSS verification"), None)
            if browser_review:
                _heading(doc, "تأیید مرورگری XSS", level=2)
                checks = browser_review.get("checks", [])
                self._paragraph(doc, f"تلاش تأیید: {len(checks)} | درخواست مرورگری: {browser_review.get('request_count', 0)}؛ جدا از شمارش درخواست‌های HTTP بررسی اولیه.")
                if checks:
                    rows = [[str(item.get("target", "")), str(item.get("parameter", "")),
                             "تأیید اجرای نشانگر" if item.get("script_execution_verified") else "تأیید نشد / نامشخص",
                             str(item.get("browser_version", "اجرا نشده"))] for item in checks]
                    _add_table(doc, ["مسیر", "پارامتر", "نتیجه", "نسخه مرورگر"], rows, [2.6, 0.9, 1.8, 1.5], font_size=8)
                self._paragraph(doc, "فقط نامزدهای بازتاب HTML در اسکن عمیق و تا سقف تعیین‌شده بررسی می‌شوند. تأیید نشدن اجرا به معنی نبود XSS نیست؛ پوشش XSS ذخیره‌شده، نشست احرازشده و تمام مسیرهای DOM انجام نشده است.")
            coverage = next((entry for entry in scan_run.execution_log if entry.get("operation") == "Assessment coverage"), None)
            review = next((entry for entry in scan_run.execution_log if entry.get("operation") == "Session/access review"), None)
            if coverage or review:
                _heading(doc, "پوشش نشست و کنترل دسترسی", level=2)
                if review:
                    self._paragraph(doc, f"وضعیت: {review.get('state', 'نامشخص')} | درخواست GET واقعی این بخش: {review.get('request_count', 0)}. فقط مسیرها و انتظارهای تعیین‌شده بررسی شدند؛ نشانه محتوا، متن بدنه و اعتبارنامه ذخیره نشده‌اند.")
                    rows = [[str(item.get('url', '')), str(item.get('identity', '')), str(item.get('check', '')),
                             str(item.get('http_status', '')), str(item.get('outcome', ''))] for item in review.get('checks', [])]
                    if rows:
                        _add_table(doc, ["مسیر", "هویت", "بررسی", "HTTP", "نتیجه"], rows, [2.3, 0.8, 1.9, 0.6, 1.2], font_size=8)
                    self._paragraph(doc, "PASS فقط نتیجه انتظار تعریف‌شده است؛ INCONCLUSIVE اثبات دسترسی امن نیست. نقش حساب و اثر آسیب‌پذیری باید بازبینی شوند.")
                else:
                    self._paragraph(doc, "بررسی نشست و دسترسی اجرا نشده است؛ نشست معتبر و مسیرهای خواندنی صریح لازم‌اند. این بخش تأیید امنیت دسترسی نیست.")
                self._paragraph(doc, "ورود، MFA، بازیابی رمز، گردش‌کارهای تغییردهنده مانند خرید و پرداخت، مسیرهای تعریف‌نشده و اثر کامل نفوذ نیازمند سناریو و بازبینی جداگانه‌اند.")

        _heading(doc, "تعداد یافته‌های فعال به تفکیک شدت", level=2)
        severity_row = [[_SEVERITY_FA[key] for key in _SEVERITY_ORDER]]
        count_row = [[str(counts[key]) for key in _SEVERITY_ORDER]]
        summary_table = doc.add_table(rows=0, cols=5)
        _prepare_table(summary_table)
        severity_colors = ["9C0006", "C65911", "BF9000", "548235", "5B6573"]
        for row_index, values in enumerate(severity_row + count_row):
            cells = summary_table.add_row().cells
            for col, value in enumerate(values):
                cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_margins(cells[col], top=130, bottom=130)
                cells[col].width = Inches(1.36)
                _set_cell_shading(cells[col], severity_colors[col] if row_index == 0 else (_PALE if row_index % 2 else "FFFFFF"))
                _set_cell_text(cells[col], value, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "222222", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

        note = doc.add_paragraph()
        _set_rtl(note)
        _style_run(note.add_run("تفسیر شدت: "), bold=True, size=9)
        _style_run(note.add_run("بحرانی و زیاد نیازمند اولویت بالا هستند؛ متوسط و کم برای برنامهٔ اصلاح ثبت شده‌اند؛ اطلاعاتی به‌تنهایی آسیب‌پذیری محسوب نمی‌شود."), size=9)

        doc.add_page_break()
        _heading(doc, "فهرست یافته‌ها", level=1)
        ordered = sorted(
            findings,
            key=lambda item: (_SEVERITY_ORDER.get(_value(item.severity).upper(), 9), item.title.casefold()),
        )
        if ordered:
            index_rows = []
            for number, finding in enumerate(ordered, 1):
                severity = _value(finding.severity).upper()
                index_rows.append([
                    str(number),
                    _SEVERITY_FA.get(severity, severity),
                    _STATUS_FA.get(_value(finding.status).upper(), _value(finding.status)),
                    finding.title,
                ])
            _add_table(doc, ["ردیف", "شدت", "وضعیت", "عنوان یافته"], index_rows, [0.55, 0.8, 1.55, 3.9], font_size=8)
        else:
            self._paragraph(doc, "برای این اسکن یافته‌ای ثبت نشده است.")

        _heading(doc, "شرح تفصیلی یافته‌ها", level=1)
        for number, finding in enumerate(ordered, 1):
            self._write_finding(doc, number, finding)

        doc.add_page_break()
        _heading(doc, "روش ارزیابی و محدودیت‌ها", level=1)
        method = (
            "این گزارش فقط یافته‌ها و شواهد ذخیره‌شده در Qolam Sentinel را بازگو می‌کند. "
            "وضعیت هر مورد در جدول همان یافته آمده است؛ گزارش اسکنر به‌تنهایی تأیید مستقل محسوب نمی‌شود."
        )
        self._paragraph(doc, method)
        self._paragraph(doc, "یافته‌های واردشده از ابزارهای بیرونی ممکن است عنوان و شرح را به زبان اصلی همان ابزار نگه دارند؛ وضعیت و منبع گزارش در کنار هر یافته مشخص است.")
        if scan_run and any(name in scan_run.tool_name for name in ("Web Security Auditor", "Sentinel DAST", "Sentinel Web")):
            self._paragraph(doc, "محدودهٔ موتور وب: حداکثر ۲۵ صفحه از همان مبدأ وب، درخواست GET، بدون ارسال فرم و بدون دنبال‌کردن تغییرمسیر. هر پاسخ حداکثر تا ۱ مگابایت برای تحلیل دریافت می‌شود.")
            self._paragraph(doc, "در اسکن عمیق، برای حداکثر سه نامزد بازتاب ورودی، تأیید اجرای نشانگر بی‌خطر XSS در مرورگر تلاش می‌شود؛ وضعیت واقعی هر تلاش در شواهد آمده است. XSS ذخیره‌شده، تمام مسیرهای DOM، بهره‌برداری SQL و منطق کامل کسب‌وکار پوشش داده نمی‌شوند. بررسی نشست فقط در صورت تنظیم صریح انجام می‌شود. نبودن یافته، نبود قطعی آسیب‌پذیری را ثابت نمی‌کند.")
        if scan_run and "Nmap" in scan_run.tool_name:
            self._paragraph(doc, "این گزارش شامل مشاهدهٔ پورت و سرویس از خروجی Nmap است. باز بودن پورت به‌تنهایی به معنی وجود آسیب‌پذیری نیست.")

        _heading(doc, "وضعیت ابزارها", level=1)
        if tools_health:
            tool_rows = [
                [
                    tool.name,
                    tool.availability_status,
                    tool.version or "نامشخص",
                    tool.status_message or "پیام وضعیت ثبت نشده است",
                ]
                for tool in tools_health
            ]
            _add_table(doc, ["ابزار", "وضعیت", "نسخه", "توضیح"], tool_rows, [1.25, 1.0, 1.2, 3.35], font_size=8)
        else:
            self._paragraph(doc, "اطلاعات سلامت ابزارها در زمان تهیهٔ گزارش موجود نبود.")

        doc.core_properties.title = report_title or "گزارش ارزیابی امنیتی"
        doc.core_properties.subject = f"گزارش امنیتی برای {target}"
        doc.core_properties.author = "III85III QOLAM SENTINEL"
        filename = f"Sentinel_Report_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.docx"
        out_path = self.output_dir / filename
        doc.save(out_path)
        return out_path

    @staticmethod
    def _paragraph(doc, text: str) -> None:
        paragraph = doc.add_paragraph()
        _set_rtl(paragraph)
        _style_run(paragraph.add_run(text), size=10)

    def _write_finding(self, doc, number: int, finding: Finding) -> None:
        severity = _value(finding.severity).upper()
        status = _value(finding.status).upper()
        confidence = _value(finding.confidence).upper()
        title_fa, plain_explanation, remediation_fa = _plain_finding_text(finding)
        _heading(doc, f"{number}. {_SEVERITY_FA.get(severity, severity)} — {title_fa}", level=2)
        details = [
            ["شناسهٔ یافته", finding.id],
            ["زمان ثبت", _local_time(finding.timestamp)],
            ["شدت", f"{_SEVERITY_FA.get(severity, severity)} — {_SEVERITY_MEANING.get(severity, 'شدت طبق گزارش منبع ثبت شده است.')}"],
            ["مبنای شدت", finding.severity_source or "در گزارش منبع مشخص نشده است"],
            ["وضعیت بررسی", _STATUS_FA.get(status, status)],
            ["اطمینان گزارش", f"{_CONFIDENCE_FA.get(confidence, confidence)} — {_CONFIDENCE_MEANING.get(confidence, 'سطحی که ابزار منبع اعلام کرده است.')}"],
            ["محل / مؤلفه", finding.affected_component or finding.target or "در گزارش منبع مشخص نشده است"],
            ["شناسهٔ استاندارد", "  |  ".join(item for item in (finding.cve_id, finding.cwe_id) if item) or "ثبت نشده"],
            ["خلاصهٔ قابل‌فهم", plain_explanation],
            ["توضیح فنی منبع", finding.description or "توضیحی در گزارش منبع ثبت نشده است"],
            ["خلاصهٔ شواهد", finding.evidence_summary or "خلاصهٔ جداگانه ثبت نشده است"],
            ["دلیل اطمینان اعلام‌شده", finding.confidence_justification or "در گزارش منبع ذکر نشده است"],
            ["شاهد فنی خام با حذف داده‌های حساس", _safe_evidence(finding)],
            ["راهکار اصلاح", remediation_fa],
            ["راهکار ثبت‌شدهٔ ابزار", finding.remediation or "در گزارش منبع ثبت نشده است"],
            ["ابزار و نسخه", f"{finding.source_tool}  |  {finding.tool_version}"],
        ]
        if finding.cwe_id == "CWE-79":
            try:
                evidence = json.loads(finding.evidence_raw)
            except (ValueError, TypeError):
                evidence = {}
            if isinstance(evidence, dict) and evidence.get("verification_method") == "inert_marker_reflection":
                details.append(["تأیید اجرای XSS", "در این بررسی، اجرای کد در مرورگر تأیید نشده است؛ شاهد موجود، بازتاب نشانگر بدون اجرای اسکریپت است."])
            if isinstance(evidence, dict) and evidence.get("verification_method") == "browser_execution_probe":
                proof = evidence.get("browser_verification", {})
                if (evidence.get("script_execution_verified") is True and proof.get("script_execution_verified") is True
                        and proof.get("proof_expected") and proof.get("proof_expected") == proof.get("proof_observed")):
                    details.append(["تأیید اجرای XSS", "اجرای نشانگر بی‌خطر در سند اصلی مرورگر مشاهده شد؛ مقدار محاسبه‌شده با مقدار مورد انتظار برابر بود. اثر روی حساب کاربران بررسی نشده است."])
                    details.append(["مرورگر تأییدکننده", str(proof.get("browser_name", "")) + " " + str(proof.get("browser_version", ""))])
                    details.append(["پارامتر آسیب‌پذیر", str(proof.get("parameter", ""))])
                    details.append(["ورودی بی‌خطر آزمون", str(proof.get("payload", ""))])
                    details.append(["مقدار مورد انتظار / مشاهده‌شده", str(proof.get("proof_expected", "")) + " / " + str(proof.get("proof_observed", ""))])
                    details.append(["بازآزمایی", "در مسیر و پارامتر ثبت‌شده، ورودی آزمون را فقط در محدودهٔ مجاز قرار دهید و اجرای همان محاسبه را در سند اصلی بررسی کنید. مقادیر خصوصی دیگر پارامترها ذخیره نشده‌اند و باید توسط ارزیاب تکمیل شوند."])
                else:
                    details.append(["تأیید اجرای XSS", "شواهد کافی برای تأیید اجرای کد در مرورگر ثبت نشده است."])
        if title_fa != finding.title:
            details.insert(5, ["عنوان اصلی ابزار", finding.title])
        table = _add_table(doc, ["بخش", "جزئیات"], details, [1.45, 5.35], font_size=8)
        if table.rows:
            for cell in table.rows[0].cells:
                _set_cell_shading(cell, _NAVY)
            for row_index, row in enumerate(table.rows[1:], 1):
                _set_cell_shading(row.cells[0], _BLUE)
                _set_cell_text(row.cells[0], details[row_index - 1][0], bold=True, size=8)
                _set_cell_text(row.cells[1], details[row_index - 1][1], size=8)
        note = doc.add_paragraph()
        _set_rtl(note)
        _style_run(note.add_run("نکته: "), bold=True, size=8, color="667085")
        _style_run(note.add_run("شدت و شرح فنی از منبع ثبت‌شده آمده‌اند؛ موردی با وضعیت گزارش‌شده یا احتمالی، پیش از اقدام پرخطر باید دستی بازبینی شود."), size=8, color="667085")


class JSONExporter:
    """Exports structured findings and scans to standardized JSON."""

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or REPORTS_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_json(
        self,
        scan_run: Optional[ScanRun],
        findings: List[Finding],
        tools_health: List[ToolHealthInfo],
    ) -> Path:
        data = {
            "platform": "III85III QOLAM SENTINEL",
            "export_timestamp": datetime.now().astimezone().isoformat(),
            "scan": scan_run.to_dict() if scan_run else None,
            "tool_health": [t.to_dict() for t in tools_health],
            "total_findings": len(findings),
            "findings": [f.to_dict() for f in findings],
        }
        filename = f"Sentinel_Export_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.json"
        out_path = self.output_dir / filename
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(sanitize_data(data), f, ensure_ascii=False, indent=2)
        return out_path
