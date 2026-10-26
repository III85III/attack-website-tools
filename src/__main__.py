"""
III85III OBSIDIAN — Application Entry Point
Evidence-Driven Cybersecurity Platform for Windows
"""
from __future__ import annotations

import os
import sys
import ipaddress
import json
from pathlib import Path
from urllib.parse import urlsplit

# Ensure project root is in sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication

from src.core.config import APP_VERSION
from src.core.database import DatabaseManager
from src.ui.main_window import MainWindow


def _verify_startup(output_path: str) -> int:
    """Release QA with a disposable database, no targets and no network scans."""
    import tempfile
    from PyQt6.QtCore import QTimer
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    app = QApplication.instance() or QApplication([])
    from PyQt6.QtGui import QFont, QFontDatabase
    font_path = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "Fonts" / "segoeui.ttf"
    if font_path.is_file():
        QFontDatabase.addApplicationFont(str(font_path))
        app.setFont(QFont("Segoe UI", 10))
    app.setApplicationName("OBSIDIAN")
    app.setOrganizationName("III85III")
    old_database = DatabaseManager._instance
    with tempfile.TemporaryDirectory(prefix="obsidian-startup-") as folder:
        db = DatabaseManager(Path(folder) / "qa.db")
        DatabaseManager._instance = db
        window = None
        try:
            window = MainWindow()
            window.show()
            app.processEvents()
            destination = Path(output_path).resolve()
            destination.parent.mkdir(parents=True, exist_ok=True)
            screenshot = destination.with_suffix(".png")
            screenshot_ok = window.grab().save(str(screenshot))
            widget = window._tabs[0]
            result = {"application_version": APP_VERSION, "window_title": window.windowTitle(),
                "scan_button": widget.scan_btn.text(), "target_input_empty": not widget.target_input.text(),
                "history_isolated": True, "screenshot_saved": screenshot_ok}
            destination.write_text(json.dumps(result, indent=2), encoding="utf-8")
            return 0 if screenshot_ok and "OBSIDIAN" in window.windowTitle() and result["target_input_empty"] else 1
        finally:
            if window:
                window.close()
            db.close()
            DatabaseManager._instance = old_database


def _verify_local_scanner(target_url: str, output_path: str, deep_scan: bool = False) -> int:
    """Run a real scanner smoke check against loopback only, for release QA."""
    try:
        parsed = urlsplit(target_url)
        host = parsed.hostname or ""
        try:
            loopback = ipaddress.ip_address(host).is_loopback
        except ValueError:
            loopback = host.lower() == "localhost"
        if not loopback or parsed.scheme.lower() not in {"http", "https"}:
            raise ValueError("The packaged scanner self-check accepts loopback HTTP/S targets only.")

        from src.adapters.web_auditor import WebAuditorAdapter
        from src.core.models import ScanRun

        scan = ScanRun(tool_name="Packaged OBSIDIAN DAST self-check", target=target_url, raw_command="")
        scan, findings = WebAuditorAdapter().run_scan(target_url, scan, {"deep_scan": deep_scan})
        import tempfile
        from src.reporting.export_service import export_assessment
        from src.core.privacy import sanitize_data

        persisted_ids = set()
        exported_ids = set()
        with tempfile.TemporaryDirectory(prefix="sentinel-self-check-") as temporary:
            database = DatabaseManager(Path(temporary) / "verification.db")
            try:
                database.save_scan(scan)
                for finding in findings:
                    database.save_finding(finding)
                database.save_scan(scan)
                persisted_ids = {item.id for item in database.list_findings(scan_id=scan.id, limit=None)}
                if scan.status not in {"RUNNING", "QUEUED"}:
                    report = export_assessment(database, scan.id, "json", Path(temporary))
                    exported_ids = {item["id"] for item in json.loads(report.read_text(encoding="utf-8"))["findings"]}
            finally:
                database.close()
        expected_ids = {finding.id for finding in findings}
        persistence_verified = persisted_ids == expected_ids
        export_verified = exported_ids == expected_ids
        payload = {
            "application_version": APP_VERSION,
            "scan": scan.to_dict(),
            "findings": [finding.to_dict() for finding in findings],
            "persistence_verified": persistence_verified,
            "export_verified": export_verified,
        }
        destination = Path(output_path).resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(json.dumps(sanitize_data(payload), ensure_ascii=False, indent=2), encoding="utf-8")
        return 0 if (scan.status == "COMPLETED" and scan.request_count > 0
                     and scan.response_status is not None and persistence_verified and export_verified) else 1
    except Exception as exc:
        try:
            Path(output_path).resolve().write_text(
                json.dumps({"application_version": APP_VERSION, "status": "FAILED", "error_type": type(exc).__name__}, indent=2),
                encoding="utf-8",
            )
        except OSError:
            pass
        return 1


def _verify_local_session_review(target_url: str, plan_path: str, output_path: str) -> int:
    """Exercise supplied synthetic sessions against loopback only for release QA."""
    try:
        parsed = urlsplit(target_url)
        try:
            local = ipaddress.ip_address(parsed.hostname or "").is_loopback
        except ValueError:
            local = (parsed.hostname or "").lower() == "localhost"
        if not local or parsed.scheme not in {"http", "https"}:
            raise ValueError("Session release verification accepts loopback HTTP/S only.")
        with Path(plan_path).open("rb") as source:
            raw = source.read(65_537)
        if len(raw) > 65_536:
            raise ValueError("Verification plan exceeds the size limit.")
        from src.adapters.session_review import SessionReviewAdapter
        from src.core.models import ScanRun
        from src.core.privacy import sanitize_data
        from src.reporting.export_service import export_assessment
        import tempfile
        scan = ScanRun(tool_name=SessionReviewAdapter.name, target=target_url, raw_command="Loopback session QA")
        scan, findings = SessionReviewAdapter().run_scan(target_url, scan, json.loads(raw))
        with tempfile.TemporaryDirectory(prefix="sentinel-session-qa-") as temporary:
            database = DatabaseManager(Path(temporary) / "verification.db")
            try:
                database.save_import(scan, findings)
                stored = {item.id for item in database.list_findings(scan_id=scan.id, limit=None)}
                report = export_assessment(database, scan.id, "json", Path(temporary))
                exported = {item["id"] for item in json.loads(report.read_text(encoding="utf-8"))["findings"]}
            finally:
                database.close()
        expected = {item.id for item in findings}
        payload = {"application_version": APP_VERSION, "scan": scan.to_dict(),
                   "findings": [item.to_dict() for item in findings],
                   "persistence_verified": stored == expected, "export_verified": exported == expected}
        Path(output_path).write_text(json.dumps(sanitize_data(payload), indent=2), encoding="utf-8")
        return 0 if scan.status == "COMPLETED" and scan.request_count > 0 and stored == exported == expected else 1
    except Exception as exc:
        Path(output_path).write_text(json.dumps({"application_version": APP_VERSION, "status": "FAILED", "error_type": type(exc).__name__}), encoding="utf-8")
        return 1


def _verify_local_api(target_url: str, spec_path: str, output_path: str) -> int:
    """Exercise the packaged API engine and actual Word export on loopback only."""
    try:
        parsed = urlsplit(target_url)
        try:
            local = ipaddress.ip_address(parsed.hostname or "").is_loopback
        except ValueError:
            local = (parsed.hostname or "").lower() == "localhost"
        if not local or parsed.scheme not in {"http", "https"}:
            raise ValueError("API release verification accepts loopback HTTP/S only.")
        import tempfile
        from zipfile import ZipFile
        from docx import Document
        from src.adapters.api_review import APIReviewAdapter, import_spec_file
        from src.core.models import ScanRun
        from src.core.privacy import sanitize_data
        from src.reporting.export_service import export_assessment
        plan = import_spec_file(target_url, spec_path)
        scan, findings = APIReviewAdapter().run_scan(target_url, ScanRun(tool_name=APIReviewAdapter.name, target=target_url, raw_command=""), plan)
        with tempfile.TemporaryDirectory(prefix="sentinel-api-qa-") as temporary:
            database = DatabaseManager(Path(temporary) / "verification.db")
            try:
                database.save_import(scan, findings)
                json_path = export_assessment(database, scan.id, "json", Path(temporary))
                word_path = export_assessment(database, scan.id, "docx", Path(temporary))
                persisted = {item.id for item in database.list_findings(scan_id=scan.id, limit=None)}
                exported = {item["id"] for item in json.loads(json_path.read_text(encoding="utf-8"))["findings"]}
                with ZipFile(word_path) as archive:
                    valid_docx = archive.testzip() is None
                document = Document(word_path)
                text = "\n".join([p.text for p in document.paragraphs] + [c.text for table in document.tables for row in table.rows for c in row.cells])
                word_ok = valid_docx and "پوشش مسیرهای API" in text and all(item.id in text for item in findings)
            finally:
                database.close()
        expected = {item.id for item in findings}
        payload = {"application_version": APP_VERSION, "scan": scan.to_dict(), "findings": [item.to_dict() for item in findings],
                   "persistence_verified": persisted == expected, "export_verified": exported == expected,
                   "word_export_verified": word_ok, "word_visual_review": False}
        Path(output_path).write_text(json.dumps(sanitize_data(payload), indent=2), encoding="utf-8")
        return 0 if scan.request_count > 0 and scan.status in {"COMPLETED", "PARTIAL"} and persisted == exported == expected and word_ok else 1
    except Exception as exc:
        Path(output_path).write_text(json.dumps({"application_version": APP_VERSION, "status": "FAILED", "error_type": type(exc).__name__}), encoding="utf-8")
        return 1


def main():
    if len(sys.argv) == 3 and sys.argv[1] == "--verify-startup":
        return _verify_startup(sys.argv[2])
    if len(sys.argv) == 5 and sys.argv[1] == "--verify-local-api":
        return _verify_local_api(sys.argv[2], sys.argv[3], sys.argv[4])
    if len(sys.argv) == 4 and sys.argv[1] == "--verify-local-browser":
        return _verify_local_scanner(sys.argv[2], sys.argv[3], deep_scan=True)
    if len(sys.argv) == 4 and sys.argv[1] == "--verify-local-scanner":
        return _verify_local_scanner(sys.argv[2], sys.argv[3])
    if len(sys.argv) == 5 and sys.argv[1] == "--verify-local-session":
        return _verify_local_session_review(sys.argv[2], sys.argv[3], sys.argv[4])

    app = QApplication(sys.argv)
    app.setApplicationName("OBSIDIAN")
    app.setOrganizationName("III85III")
    app.setApplicationVersion(APP_VERSION)

    from src.ui.usage_notice_dialog import ensure_usage_notice
    if not ensure_usage_notice():
        return 0

    # Initialize the workspace only after the responsible-use notice.
    db = DatabaseManager.get_instance()

    window = MainWindow()
    window.show()

    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
