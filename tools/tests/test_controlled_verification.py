"""Actual button-to-network-to-persisted-Word verification and safety controls."""
import json
import os
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlsplit, parse_qs
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import pytest
from docx import Document
from PyQt6.QtCore import QEventLoop, QTimer
from PyQt6.QtWidgets import QApplication, QMessageBox
from src.adapters.web_auditor import WebAuditorAdapter
from src.adapters.resilience_review import run_resilience, validate_resilience
from src.core.application_map import origin
from src.core.database import DatabaseManager
from src.core.models import ScanRun, ToolHealthInfo
from src.reporting.export_service import export_assessment
from src.ui.tabs import WebSecurityTab

@pytest.fixture
def target():
    events = []
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            events.append(self.path)
            parsed = urlsplit(self.path)
            value = parse_qs(parsed.query).get('next', [''])[0]
            status, body, destination = 200, '<a href="/redirect?next=/welcome">Navigation</a>', None
            if parsed.path == '/redirect' and value:
                status, destination = 302, value
                body = 'Redirect response'
            elif parsed.path == '/fixed':
                status, destination, body = 302, 'https://fixed.invalid/login', 'Fixed destination'
            elif parsed.path == '/backoff':
                status, body = 429, 'Back off'
            elif parsed.path not in {'/', '/secure'}:
                status, body = 404, 'Not found'
            data = body.encode()
            self.send_response(status)
            self.send_header('Content-Type', 'text/html')
            self.send_header('Content-Length', str(len(data)))
            if destination: self.send_header('Location', destination)
            self.end_headers(); self.wfile.write(data)
        def log_message(self, *_args): pass
    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    worker = threading.Thread(target=server.serve_forever, daemon=True); worker.start()
    try: yield f'http://127.0.0.1:{server.server_port}', events
    finally: server.shutdown(); server.server_close(); worker.join(3)

@pytest.mark.parametrize('path,verified', [('/redirect', True), ('/secure', False), ('/fixed', False)])
def test_controlled_redirect_requires_challenges_and_controls(target, path, verified):
    base, events = target
    from urllib.parse import urlparse, parse_qsl
    parsed = urlparse(base + path + '?next=/welcome')
    adapter = WebAuditorAdapter(); adapter._authorized_origin = origin(base)
    findings = adapter._verify_open_redirect(parsed, parse_qsl(parsed.query), 'next', 'fixture', base + path)
    assert bool(findings) == verified
    assert len(events) == (4 if verified else 3)
    if findings:
        assert findings[0].status.value == 'VALIDATED'
        proof = json.loads(findings[0].evidence_raw)
        assert proof['outcome'] == 'VERIFIED' and len(proof['observations']) == 4
        assert proof['observations'][2]['request']['synthetic_input'] != proof['observations'][3]['request']['synthetic_input']


def test_actual_scan_button_reaches_network_store_ui_and_word(target, tmp_path, monkeypatch):
    base, events = target
    app = QApplication.instance() or QApplication([])
    database = DatabaseManager(tmp_path / 'click.db')
    monkeypatch.setattr(DatabaseManager, 'get_instance', classmethod(lambda cls, *_a, **_kw: database))
    monkeypatch.setattr('src.ui.tabs.ZapAdapter.detect', lambda self: ToolHealthInfo(name='ZAP', installed=False))
    monkeypatch.setattr('src.ui.tabs.NmapAdapter.detect', lambda self: ToolHealthInfo(name='Nmap', installed=False))
    monkeypatch.setattr('src.ui.tabs.write_execution_log', lambda scan: tmp_path / 'run.json')
    monkeypatch.setattr(QMessageBox, 'question', lambda *args, **kwargs: QMessageBox.StandardButton.Yes)
    widget = WebSecurityTab()
    try:
        widget.target_input.setText(base)
        from src.core.traffic_inventory import import_har
        capture = tmp_path / 'overlapping.har'
        capture.write_text(json.dumps({'log': {'entries': [{'request': {'url': base + '/redirect?next=/captured', 'method': 'GET'}, 'response': {'status': 200}}]}}))
        widget.traffic_plan = import_har(capture, base)
        widget.traffic_probes_approved = True
        widget.resilience_plan = approved(base)
        monkeypatch.setattr('src.adapters.resilience_review.psutil.virtual_memory', lambda: type('Memory', (), {'percent': 10})())
        widget.scan_btn.click()  # Actual user callback creates and starts the real QThread.
        assert widget.web_worker is not None
        results = []
        loop = QEventLoop()
        widget.web_worker.finished_signal.connect(lambda scan, findings: (results.append((scan, findings)), loop.quit()))
        QTimer.singleShot(30000, loop.quit); loop.exec()
        assert widget.web_worker.wait(5000)
        app.processEvents()
        assert results
        scan, findings = results[0]
        assert scan.request_count == len(events) > 4
        confirmed = [f for f in findings if f.cwe_id == 'CWE-601' and f.status.value == 'VALIDATED']
        assert len(confirmed) == 1
        coverage = next(e for e in scan.execution_log if e.get('operation') == 'Route verification coverage')
        assert coverage['security_checked_inputs'] >= 1
        assert sum(row['verified_findings'] for row in coverage['routes']) == 1
        saved = database.list_findings(scan_id=scan.id, limit=None)
        assert {f.id for f in saved} == {f.id for f in findings}
        assert widget.findings_table.rowCount() == len(findings)
        path = export_assessment(database, scan.id, 'docx', tmp_path)
        document = Document(path)
        text = '\n'.join([p.text for p in document.paragraphs] + [c.text for t in document.tables for row in t.rows for c in row.cells])
        assert confirmed[0].id in text and 'controlled_open_redirect_v2' in text
        assert 'Controlled resilience sample' in text
        assert any(e.get('operation') == 'Controlled resilience sample' and e['request_count'] == 2 for e in scan.execution_log)
        exported = json.loads(export_assessment(database, scan.id, 'json', tmp_path).read_text())
        assert any(f['id'] == confirmed[0].id for f in exported['findings'])
    finally:
        if widget.web_worker.isRunning(): widget.web_worker.requestInterruption(); widget.web_worker.wait(15000)
        widget.close(); database.close()


def approved(base, **kw):
    return {'owned_read_only_approved': True, 'approved_origin': origin(base), 'requests': 2, 'rate': 2, **kw}


def test_resilience_limits_real_requests_and_stops_on_backpressure(target, monkeypatch):
    base, events = target
    monkeypatch.setattr('src.adapters.resilience_review.psutil.virtual_memory', lambda: type('Memory', (), {'percent': 10})())
    sample = run_resilience(base + '/secure', approved(base))
    assert sample['request_count'] == len(events) == 2 and sample['state'] == 'COMPLETED'
    sample = run_resilience(base + '/backoff', approved(base))
    assert sample['request_count'] == 1 and sample['stop_reason'] == 'Target backpressure response.'
    before = len(events)
    sample = run_resilience(base, approved(base), lambda: True)
    assert sample['request_count'] == 0 and len(events) == before

@pytest.mark.parametrize('mutation', ['approval', 'origin', 'budget', 'method', 'nan'])
def test_resilience_rejects_unapproved_or_prohibited_plans_before_network(target, mutation):
    base, events = target
    plan = approved(base); target_url = base
    if mutation == 'approval': plan['owned_read_only_approved'] = False
    if mutation == 'origin': plan['approved_origin'] = origin('https://elsewhere.invalid')
    if mutation == 'budget': plan['requests'] = 1000
    if mutation == 'method': target_url += '/delete'
    if mutation == 'nan': plan['rate'] = float('nan')
    with pytest.raises(ValueError): run_resilience(target_url, plan)
    assert not events


def test_network_layer_blocks_out_of_scope_and_state_changing_methods(target):
    base, events = target
    import requests
    adapter = WebAuditorAdapter(); adapter._authorized_origin = origin(base)
    for url, method in [('https://outside.invalid/', 'GET'), (base, 'POST'), (base + '/delete', 'GET')]:
        with pytest.raises(requests.RequestException): adapter._bounded_get(url, timeout=2, method=method)
    assert not events


def test_resilience_duration_and_client_resource_stop(target, monkeypatch):
    base, events = target
    monkeypatch.setattr('src.adapters.resilience_review.psutil.virtual_memory', lambda: type('Memory', (), {'percent': 99})())
    sample = run_resilience(base, approved(base))
    assert sample['request_count'] == 0 and 'memory threshold' in sample['stop_reason']
    monkeypatch.setattr('src.adapters.resilience_review.psutil.virtual_memory', lambda: type('Memory', (), {'percent': 10})())
    sample = run_resilience(base, approved(base, rate=.2, duration_seconds=1))
    assert sample['request_count'] == len(events) == 1
    assert sample['state'] == 'INCONCLUSIVE' and sample['stop_reason'] == 'Duration admission limit reached.'
    assert sample['duration_ms'] < 2500
