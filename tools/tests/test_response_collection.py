import json
from unittest.mock import MagicMock

import pytest
import requests
from docx import Document

from src.adapters import web_auditor as module
from src.core.models import ScanRun
from src.reporting.docx_report import JSONExporter, ReportGenerator


def response(body=b"<html>Local fixture</html>", broken=False):
    result = requests.Response()
    result.status_code = 200
    result.headers["Content-Type"] = "text/html"
    result.encoding = "utf-8"
    result.raw = MagicMock()
    def chunks(**kwargs):
        yield body
        if broken:
            raise requests.exceptions.ChunkedEncodingError("token=FIXTURE_PRIVATE")
    result.iter_content = chunks
    return result


def run(adapter):
    scan = ScanRun(tool_name="Fixture", target="http://127.0.0.1/", raw_command="")
    return adapter.run_scan(scan.target, scan)


def test_initial_body_failure_is_failed_despite_http_200(monkeypatch):
    adapter = module.WebAuditorAdapter()
    item = response(broken=True)
    monkeypatch.setattr(adapter.session, "get", lambda *a, **kw: item)
    scan, findings = run(adapter)
    assert scan.response_status == 200 and scan.request_count == 1
    assert scan.status == "FAILED" and scan.exit_code != 0
    assert "body" in scan.completion_reason
    assert scan.execution_log[-1]["response_body_read_errors"] == 1
    assert scan.execution_log[-1]["request_error_types"] == ["ChunkedEncodingError"]
    assert "FIXTURE_PRIVATE" not in json.dumps(scan.execution_log)
    item.raw.close.assert_called_once()


@pytest.mark.parametrize("body", [b"12345678", b"1234"])
def test_body_cap_is_partial_and_visible_in_word_json(tmp_path, monkeypatch, body):
    monkeypatch.setattr(module, "MAX_RESPONSE_BYTES", 4)
    adapter = module.WebAuditorAdapter()
    responses = []
    def get(*args, **kwargs):
        value = response(body)
        responses.append(value)
        return value
    monkeypatch.setattr(adapter.session, "get", get)
    scan, findings = run(adapter)
    assert scan.status == "PARTIAL"
    assert "limit hit(s)" in scan.completion_reason
    assert scan.execution_log[-1]["response_body_limit_bytes"] == 4
    assert scan.execution_log[-1]["response_body_limit_hits"] == len(responses)
    assert all(len(item.content) == 4 for item in responses)
    payload = JSONExporter(tmp_path).export_json(scan, findings, [])
    assert json.loads(payload.read_text(encoding="utf-8"))["scan"]["status"] == "PARTIAL"
    word = Document(ReportGenerator(tmp_path).generate_docx(scan, findings, []))
    text = "\n".join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
    assert "limit hit(s)" in text


def test_followup_body_failures_are_partial_and_counted_once(monkeypatch):
    adapter = module.WebAuditorAdapter()
    calls = []
    def get(*args, **kwargs):
        calls.append(args)
        return response(broken=len(calls) > 1)
    monkeypatch.setattr(adapter.session, "get", get)
    scan, findings = run(adapter)
    assert scan.status == "PARTIAL" and scan.response_status == 200
    log = scan.execution_log[-1]
    assert log["response_body_read_errors"] == len(calls) - 1
    assert len(log["request_error_types"]) == len(calls) - 1


def test_no_verified_response_cannot_be_success(monkeypatch):
    adapter = module.WebAuditorAdapter()
    monkeypatch.setattr(adapter, "audit_url", lambda *args, **kwargs: [])
    scan, findings = run(adapter)
    assert scan.status == "FAILED" and scan.request_count == 0


def test_pipeline_preserves_partial_status_and_does_not_start_optional_engines(tmp_path, monkeypatch):
    from PyQt6.QtWidgets import QApplication
    from src.ui.tabs import WebAuditWorker
    from src.core.database import DatabaseManager
    app = QApplication.instance() or QApplication([])
    database = DatabaseManager(tmp_path / "worker.db")
    monkeypatch.setattr(DatabaseManager, "get_instance", classmethod(lambda cls: database))
    monkeypatch.setattr("src.ui.tabs.write_execution_log", lambda scan: tmp_path / "fixture.json")
    worker = WebAuditWorker("http://127.0.0.1/")
    def limited(target, scan, options):
        scan.status = "PARTIAL"
        scan.response_status = 200
        scan.completion_reason = "Fixture body limit reached"
        return scan, []
    monkeypatch.setattr(worker.adapter, "run_scan", limited)
    monkeypatch.setattr("src.ui.tabs.ZapAdapter.detect", lambda *a: pytest.fail("Optional engine should not start"))
    monkeypatch.setattr("src.ui.tabs.NmapAdapter.detect", lambda *a: pytest.fail("Optional engine should not start"))
    results = []
    worker.finished_signal.connect(lambda scan, findings: results.append(scan))
    worker.run()
    assert results[0].status == "PARTIAL"
    assert results[0].execution_log[-2]["state"] == "PARTIAL"
    assert database.get_scan(results[0].id).status == "PARTIAL"
    database.close()
