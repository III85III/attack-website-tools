"""Real local HAR-to-check-to-store-to-report corpus and negative controls."""
import json
import os
import threading
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import pytest
from docx import Document
from PyQt6.QtWidgets import QApplication

from src.adapters.web_auditor import WebAuditorAdapter
from src.core.database import DatabaseManager
from src.core.models import Finding, FindingConfidence, FindingSeverity, FindingStatus, ScanRun, ToolHealthInfo
from src.core.traffic_inventory import import_har, validate_traffic_plan, synthetic_input_points
from src.reporting.export_service import export_assessment

PRIVATE_QUERY = "CAPTURED_PRIVATE_SEARCH_VALUE"
PRIVATE_COOKIE = "CAPTURED_PRIVATE_SESSION_VALUE"


@pytest.fixture
def target():
    observed = []
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            parsed = urlsplit(self.path)
            observed.append({"url": self.path, "cookie": self.headers.get("Cookie"), "authorization": self.headers.get("Authorization")})
            status, body, kind = 200, "<html>Synthetic root, no hidden links</html>", "text/html"
            if parsed.path in {"/unlinked", "/secure", "/always-error"}:
                value = parse_qs(parsed.query).get("search", [""])[0]
                kind = "text/plain"
                body = "sqlite3.OperationalError: synthetic local fixture" if parsed.path == "/always-error" or (parsed.path == "/unlinked" and value == "'") else "Synthetic search response"
            elif parsed.path != "/":
                status, body = 404, "Not found"
            encoded = body.encode()
            self.send_response(status); self.send_header("Content-Type", kind)
            self.send_header("Content-Length", str(len(encoded))); self.end_headers(); self.wfile.write(encoded)
        def do_POST(self):
            pytest.fail("No state-changing request is allowed in this fixture")
        def log_message(self, *_args):pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    try:yield f"http://127.0.0.1:{server.server_port}", observed
    finally:server.shutdown(); server.server_close(); thread.join(3)


def har_entry(url, method="GET", authenticated=False):
    return {"request": {"url": url, "method": method,
        "headers": [{"name": "Authorization", "value": PRIVATE_COOKIE}] if authenticated else [],
        "cookies": [{"name": "sid", "value": PRIVATE_COOKIE}] if authenticated else [],
        "postData": {"text": PRIVATE_COOKIE}}, "response": {"status": 200, "content": {"text": PRIVATE_COOKIE}}}


def plan(tmp_path, target, path="/unlinked", extra=None):
    entries = [har_entry(target + path + "?search=" + PRIVATE_QUERY)] + (extra or [])
    file = tmp_path / "capture.har"; file.write_text(json.dumps({"log": {"version": "1.2", "entries": entries}}), encoding="utf-8")
    return import_har(file, target)


def test_har_import_is_offline_scoped_and_redacted(target, tmp_path):
    base, observed = target
    result = plan(tmp_path, base, extra=[har_entry("https://outside.invalid/private?token=PRIVATE"), har_entry(base + "/account?user_id=7", authenticated=True), har_entry(base + "/write", "POST")])
    assert observed == []
    assert len(result["routes"]) == 3
    text = json.dumps(result)
    assert PRIVATE_QUERY not in text and PRIVATE_COOKIE not in text and "outside.invalid" not in text
    assert result["excluded"]["outside_scope_or_invalid"] == 1
    assert len(synthetic_input_points(result)) == 1
    assert all(not r["probe_eligible"] for r in result["routes"][1:])
    assert validate_traffic_plan(base, result)["routes"] == result["routes"]


@pytest.mark.parametrize("path,expected", [("/unlinked", True), ("/secure", False), ("/always-error", False)])
def test_labeled_har_assisted_detection_and_negative_controls(target, tmp_path, path, expected):
    base, observed = target
    config = plan(tmp_path, base, path)
    scan, findings = WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"deep_scan": True, "traffic_plan": config, "traffic_probes_approved": True})
    errors = [f for f in findings if f.cwe_id == "CWE-209"]
    assert bool(errors) is expected
    assert scan.request_count == len(observed)
    assert PRIVATE_QUERY not in json.dumps(observed) and PRIVATE_COOKIE not in json.dumps(observed)
    assert all(r["cookie"] is None and r["authorization"] is None for r in observed)
    coverage = next(e for e in scan.execution_log if e.get("operation") == "Route verification coverage")
    assert coverage["security_checked_discovered_routes"] == (0 if path == "/always-error" else 1)
    hypotheses = next(e for e in scan.execution_log if e.get("operation") == "Research hypotheses")["hypotheses"]
    assert len(hypotheses) == 1 and hypotheses[0]["requests_collected"] == (1 if path == "/always-error" else 4)
    assert bool(hypotheses[0]["finding_ids"]) is expected
    if errors:
        assert errors[0].status == FindingStatus.POTENTIAL and errors[0].confidence != FindingConfidence.CONFIRMED
        assert hypotheses[0]["state"] == "SUPPORTED_CANDIDATE"
    else:
        assert hypotheses[0]["state"] == ("INCONCLUSIVE" if path == "/always-error" else "NOT_REPRODUCED")


def test_discovery_only_har_sends_no_imported_route_requests(target, tmp_path):
    base, observed = target
    config = plan(tmp_path, base)
    scan, _ = WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"deep_scan": True, "traffic_plan": config})
    assert not any("/unlinked" in r["url"] for r in observed)
    coverage = next(e for e in scan.execution_log if e.get("operation") == "Route verification coverage")
    assert any("/unlinked" in r["url"] and not r["requested"] for r in coverage["routes"])


def test_authenticated_and_ownership_routes_are_never_replayed(target, tmp_path):
    base, observed = target
    file = tmp_path / "account.har"
    file.write_text(json.dumps({"log": {"entries": [har_entry(base + "/private?search=q", authenticated=True), har_entry(base + "/record?user_id=123"), har_entry(base + "/delete?search=q"), har_entry(base + "/write?search=q", "POST")]}}))
    config = import_har(file, base)
    assert not synthetic_input_points(config)
    scan, _ = WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"deep_scan": True, "traffic_plan": config, "traffic_probes_approved": True})
    assert not any(urlsplit(r["url"]).path in {"/private", "/record", "/delete", "/write"} for r in observed)


def test_har_preserves_method_and_credential_context(target, tmp_path):
    base, _ = target
    config = plan(tmp_path, base, extra=[har_entry(base + "/unlinked?search=other", authenticated=True), har_entry(base + "/unlinked?search=q", "POST")])
    scan, _ = WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"traffic_plan": config})
    coverage = next(e for e in scan.execution_log if e.get("operation") == "Route verification coverage")
    routes = [r for r in coverage["routes"] if "/unlinked" in r["url"]]
    assert len(routes) == 3
    assert {r["identity_context"] for r in routes} == {"anonymous", "captured-authenticated-not-replayed"}
    assert {r["method"] for r in routes} == {"GET", "POST"}


def test_invalid_scope_plan_fails_before_network(target, tmp_path):
    base, observed = target
    config = plan(tmp_path, base)
    config["routes"][0]["url"] = "https://outside.invalid/?search=%5Bredacted%5D"
    with pytest.raises(ValueError):
        WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"traffic_plan": config})
    assert observed == []


@pytest.mark.parametrize("payload", [{"log": {"entries": [None]}}, {"log": {"entries": [har_entry("http://x", "TRACE")]*2001}}, {"unexpected": []}])
def test_invalid_har_rejected(tmp_path, payload):
    file = tmp_path / "invalid.har"; file.write_text(json.dumps(payload))
    with pytest.raises(ValueError):import_har(file, "https://owned.invalid")


def test_har_size_limit_and_duplicate_keys(tmp_path):
    file = tmp_path / "large.har"; file.write_bytes(b" " * 5_000_001)
    with pytest.raises(ValueError, match="5 MB"):import_har(file, "https://owned.invalid")
    file.write_text('{"log":{},"log":{}}')
    with pytest.raises(ValueError):import_har(file, "https://owned.invalid")


def test_canonical_finding_key_survives_serialization_and_distinguishes_methods():
    a = Finding("Synthetic", "Fixture", "Web Security Auditor", "https://owned.invalid/search?q=secret", FindingSeverity.LOW,
        "Fixture", FindingConfidence.LOW, "Fixture", json.dumps({"method": "GET", "parameter": "q"}))
    b = replace(a, id="other", target="https://owned.invalid/search?q=other")
    assert a.deduplication_key == b.deduplication_key
    assert Finding.from_dict(a.to_dict()).deduplication_key == a.deduplication_key
    assert replace(a, evidence_raw=json.dumps({"method": "POST", "parameter": "q"})).deduplication_key != a.deduplication_key


def test_actual_worker_ui_and_word_receive_har_evidence(target, tmp_path, monkeypatch):
    from src.ui.tabs import WebAuditWorker
    from src.ui.intelligence_workspace import IntelligenceWorkspace
    from src.ui.main_window import MainWindow
    base, _ = target; app = QApplication.instance() or QApplication([])
    db = DatabaseManager(tmp_path / "assessment.db")
    monkeypatch.setattr(DatabaseManager, "get_instance", classmethod(lambda cls, *_a, **_kw: db))
    monkeypatch.setattr("src.ui.tabs.write_execution_log", lambda scan: tmp_path / "scan.json")
    monkeypatch.setattr("src.ui.tabs.ZapAdapter.detect", lambda self: ToolHealthInfo("OWASP ZAP", False, status_message="Not configured in fixture"))
    monkeypatch.setattr("src.ui.tabs.NmapAdapter.detect", lambda self: ToolHealthInfo("Nmap", False, status_message="Not configured in fixture"))
    results = []
    worker = WebAuditWorker(base, deep_scan=True, traffic_plan=plan(tmp_path, base), traffic_probes_approved=True)
    worker.finished_signal.connect(lambda scan, findings: results.append((scan, findings)))
    worker.run()
    scan, findings = results[0]
    assert any(f.cwe_id == "CWE-209" for f in findings)
    path = export_assessment(db, scan.id, "docx", tmp_path)
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "پوشش مسیرهای کشف‌شده" in text and "Reproducible SQLite error" in text
    assert PRIVATE_QUERY not in text and PRIVATE_COOKIE not in text
    workspace = IntelligenceWorkspace()
    assert workspace.routes.topLevelItemCount() > 0 and workspace.research.topLevelItemCount() == 1
    assert workspace.surface.graph_scene.items()
    workspace.research.setCurrentItem(workspace.research.topLevelItem(0)); app.processEvents()
    assert "SUPPORTED_CANDIDATE" in workspace.details.toPlainText()
    window = MainWindow()
    assert "https://t.me/parsamzmz" in window.creator_links.text()
    assert "https://github.com/III85III" in window.creator_links.text()
    window.close(); workspace.close(); db.close()


def test_legacy_assessment_does_not_invent_route_coverage(tmp_path, monkeypatch):
    from src.ui.intelligence_workspace import IntelligenceWorkspace
    app = QApplication.instance() or QApplication([])
    db = DatabaseManager(tmp_path / "legacy.db")
    monkeypatch.setattr(DatabaseManager, "get_instance", classmethod(lambda cls, *_a, **_kw: db))
    scan = ScanRun("Sentinel DAST", "https://owned.invalid", "", status="COMPLETED"); db.save_scan(scan)
    workspace = IntelligenceWorkspace()
    assert "coverage is unknown" in workspace.summary.text()
    assert workspace.routes.topLevelItemCount() == 0
    workspace.close(); db.close()
