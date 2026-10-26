"""OBSIDIAN labeled local corpus and migration/report regression checks."""
import json
import threading
from dataclasses import replace
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, parse_qsl, urlparse

import pytest
from docx import Document

from src.adapters.web_auditor import WebAuditorAdapter
from src.core.application_map import origin
from src.core.config import resolve_frozen_state
from src.core.coverage import CoverageLedger
from src.core.database import DatabaseManager
from src.core.evidence import distinct_findings
from src.core.models import Finding, FindingConfidence, FindingSeverity, FindingStatus, ScanRun
from src.reporting.export_service import export_assessment

CANARY = "OBSIDIAN-SYNTHETIC-CANARY-7d254bdc"


@pytest.fixture
def target():
    requests = []
    class Handler(BaseHTTPRequestHandler):
        quote_calls = 0
        def do_GET(self):
            parsed = urlparse(self.path)
            value = parse_qs(parsed.query).get("q", [""])[0]
            requests.append(self.path)
            body = "Normal synthetic response"
            status = 200
            if parsed.path == "/always-error" or (parsed.path == "/sql-error" and value == "'"):
                body = "sqlite3.OperationalError: synthetic fixture"
            if parsed.path == "/transient" and value == "'":
                Handler.quote_calls += 1
                if Handler.quote_calls == 1:
                    body = "sqlite3.OperationalError: transient fixture"
            if parsed.path == "/denied" and value == "'":
                status, body = 403, "sqlite3.OperationalError: denial fixture"
            if parsed.path == "/reflected":
                body = value
            if parsed.path == "/file" and value == "../obsidian-canary.txt":
                body = CANARY
            if parsed.path == "/always-canary":
                body = CANARY
            if parsed.path == "/backoff":
                status = 429
            encoded = body.encode()
            self.send_response(status)
            self.send_header("Content-Type", "text/plain")
            self.send_header("Content-Length", str(len(encoded)))
            self.end_headers()
            self.wfile.write(encoded)
        def log_message(self, *args): pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", requests
    finally:
        server.shutdown(); server.server_close(); thread.join(3)


SQL_CORPUS = [("sql-error", True), ("always-error", False), ("transient", False),
              ("denied", False), ("reflected", False), ("secure", False)]


@pytest.mark.parametrize("path,expected", SQL_CORPUS)
def test_sql_error_labeled_corpus(target, path, expected):
    base, observed = target
    adapter = WebAuditorAdapter()
    parsed = urlparse(base + "/" + path + "?q=1")
    adapter._authorized_origin = origin(base)
    rows = adapter._controlled_sql_error(parsed, parse_qsl(parsed.query), "q", "fixture", base + "/" + path)
    assert bool(rows) is expected
    assert len(observed) == (1 if path == "always-error" else 3 if path == "denied" else 4)
    assert adapter._coverage.checks[-1]["state"] == ("INCONCLUSIVE" if path in {"always-error", "denied"} else "COMPLETED")
    if rows:
        assert rows[0].cwe_id == "CWE-209"
        assert rows[0].status == FindingStatus.POTENTIAL
        assert rows[0].confidence != FindingConfidence.CONFIRMED
        proof = json.loads(rows[0].evidence_raw)
        assert len(proof["observations"]) == 4
        assert "extraction" in proof["limitations"]


@pytest.mark.parametrize("path,expected", [("file", True), ("always-canary", False), ("secure", False)])
def test_file_canary_positive_and_negative(target, path, expected):
    base, observed = target
    adapter = WebAuditorAdapter()
    parsed = urlparse(base + "/" + path + "?q=allowed.txt")
    adapter._canary_files = [{"endpoint": "/" + path, "parameter": "q", "value": "../obsidian-canary.txt", "marker": CANARY}]
    rows = adapter._verify_file_canary(parsed, parse_qsl(parsed.query), "q", "fixture", base + "/" + path)
    assert bool(rows) is expected
    assert len(observed) == 4
    if rows:
        assert CANARY not in rows[0].evidence_raw
        assert rows[0].severity == FindingSeverity.MEDIUM


def test_file_check_without_declared_canary_sends_no_requests(target):
    base, observed = target
    adapter = WebAuditorAdapter()
    parsed = urlparse(base + "/file?q=allowed.txt")
    assert adapter._verify_file_canary(parsed, parse_qsl(parsed.query), "q", None, base + "/file") == []
    assert observed == []
    assert adapter._coverage.checks[-1]["state"] == "BLOCKED"


def test_cancel_before_collection_is_not_failure_or_success(target):
    base, observed = target
    scan, _ = WebAuditorAdapter().run_scan(base, ScanRun("OBSIDIAN DAST", base, ""), {"cancel_check": lambda: True})
    assert scan.status == "CANCELLED"
    assert observed == [] and scan.request_count == 0
    ledger = next(e for e in scan.execution_log if e.get("operation") == "Verification coverage")
    assert ledger["counts"]["INCONCLUSIVE"] >= 1
    assert ledger["counts"]["BLOCKED"] >= 4


def test_new_checks_share_origin_and_budget_enforcement(target):
    base, observed = target
    adapter = WebAuditorAdapter()
    adapter._authorized_origin = origin(base)
    parsed = urlparse("http://127.0.0.1:1/sql-error?q=1")
    assert adapter._controlled_sql_error(parsed, parse_qsl(parsed.query), "q", None, parsed.geturl()) == []
    assert not observed
    parsed = urlparse(base + "/sql-error?q=1")
    adapter._request_attempts = 240
    assert adapter._controlled_sql_error(parsed, parse_qsl(parsed.query), "q", None, base) == []
    assert not observed
    assert adapter._coverage.checks[-1]["state"] == "INCONCLUSIVE"


def test_legacy_state_reused_and_new_state_named_obsidian(tmp_path):
    executable = tmp_path / "portable" / "release" / "package"
    local = tmp_path / "LocalAppData"
    assert resolve_frozen_state(executable, local) == local / "III85III" / "OBSIDIAN"
    legacy = local / "III85III" / "Qolam Sentinel"
    (legacy / "data").mkdir(parents=True)
    (legacy / "data" / "sentinel.db").write_bytes(b"preserved")
    assert resolve_frozen_state(executable, local) == legacy
    (executable.parent.parent / "data").mkdir(parents=True)
    (executable.parent.parent / "data" / "sentinel.db").write_bytes(b"portable")
    assert resolve_frozen_state(executable, local) == executable.parent.parent
    assert (legacy / "data" / "sentinel.db").read_bytes() == b"preserved"


def test_old_scan_names_and_statuses_roundtrip_without_rewriting(tmp_path):
    db = DatabaseManager(tmp_path / "sentinel.db")
    old = ScanRun("Sentinel DAST", "https://owned.invalid", "", status="COMPLETED")
    try:
        db.save_scan(old)
        assert db.get_scan(old.id).tool_name == "Sentinel DAST"
        path = export_assessment(db, old.id, "docx", tmp_path)
        assert path.name.startswith("OBSIDIAN_Report_")
        assert Document(path).core_properties.author == "III85III OBSIDIAN"
    finally: db.close()


def test_report_consolidates_exact_duplicates_and_keeps_distinct_proof(tmp_path):
    from src.reporting.docx_report import ReportGenerator
    item = Finding("Error disclosure", "Synthetic disclosure", "OBSIDIAN DAST", "https://owned.invalid", FindingSeverity.MEDIUM,
        "Synthetic error condition", FindingConfidence.HIGH, "Four response control", json.dumps({"verification_method": "fixture",
        "reproduction_steps": ["Use the synthetic record."], "limitations": "No extraction performed.", "token": "PRIVATE_CANARY"}), status=FindingStatus.POTENTIAL)
    duplicate = replace(item, id="duplicate")
    different = replace(item, id="different", target="https://owned.invalid/other")
    assert len(distinct_findings([item, duplicate, different])) == 2
    ledger = CoverageLedger(); ledger.record("Fixture check", "INCONCLUSIVE", "Synthetic response unavailable")
    scan = ScanRun("OBSIDIAN DAST", item.target, "", status="PARTIAL", execution_log=[ledger.snapshot()])
    path = ReportGenerator(tmp_path).generate_docx(scan, [item, duplicate], [])
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "PRIVATE_CANARY" not in text
    assert "1. Use the synthetic record." in text
    assert "INCONCLUSIVE" in text
    assert "هیچ یافته‌ای با وضعیت اعتبارسنجی‌شده" in text
    assert doc.core_properties.author == "III85III OBSIDIAN"


def test_form_mapping_preserves_case_defaults_and_omits_secrets():
    from src.adapters.web_auditor import _PageInventoryParser
    parser = _PageInventoryParser()
    parser.feed('''<form method="get" action="/search"><input name="SearchKey" value="blue">
        <input name="TenantID" value="7"><input name="csrfToken" value="PRIVATE">
        <input type="password" name="pwd" value="PRIVATE"><input name="disabled" disabled value="x">
        <input type="checkbox" name="Unchecked" value="yes">
        <input type="checkbox" name="Checked" checked value="yes"></form>''')
    assert parser.forms[0]["query_defaults"] == {"SearchKey": "blue", "TenantID": "7", "Checked": "yes"}
    assert parser.forms[0]["has_csrf_token"] is True


def test_target_backoff_stops_remaining_native_requests(target):
    base, observed = target
    adapter = WebAuditorAdapter()
    scan, _ = adapter.run_scan(base + "/backoff", ScanRun("OBSIDIAN DAST", base, ""), {"deep_scan": True})
    assert len(observed) == scan.request_count == 1
    assert scan.status == "PARTIAL"
    ledger = next(e for e in scan.execution_log if e.get("operation") == "Verification coverage")
    assert any(c["check"] == "Target availability" and c["state"] == "BLOCKED" for c in ledger["checks"])
    provenance = next(e for e in scan.execution_log if e.get("operation") == "Native request provenance")
    assert provenance["requests"][0]["status"] == 429


def test_request_provenance_retains_parameter_names_without_values(target):
    base, observed = target
    adapter = WebAuditorAdapter()
    adapter._request_evidence = []
    adapter._bounded_get(base + "/secure?q=PRIVATE_VALUE", timeout=2)
    assert len(observed) == 1
    assert "PRIVATE_VALUE" not in json.dumps(adapter._request_evidence)
    assert "q=" in adapter._request_evidence[0]["url"]


def test_report_export_formats_share_duplicate_consolidation(tmp_path):
    db = DatabaseManager(tmp_path / "export.db")
    scan = ScanRun("OBSIDIAN DAST", "https://owned.invalid", "", status="COMPLETED")
    item = Finding("Duplicate", "Synthetic", "fixture", scan.target, FindingSeverity.LOW,
        "Fixture", FindingConfidence.LOW, "Fixture", "No raw data", scan_id=scan.id)
    try:
        db.save_scan(scan); db.save_finding(item); db.save_finding(replace(item, id="duplicate"))
        path = export_assessment(db, scan.id, "json", tmp_path)
        assert json.loads(path.read_text(encoding="utf-8"))["total_findings"] == 1
        for fmt in ("markdown", "html", "docx"):
            assert export_assessment(db, scan.id, fmt, tmp_path).is_file()
    finally: db.close()
